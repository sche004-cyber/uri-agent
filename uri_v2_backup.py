import json
import html
import re
import subprocess
import threading
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse

# ============================================================
# URI V2 — NIT Sikkim Administrative AI Assistant
# Local document retrieval + Hermes AI
# ============================================================

BASE = Path(__file__).resolve().parent
KNOWLEDGE = BASE / "knowledge"
CACHE = BASE / "cache"
CACHE.mkdir(exist_ok=True)

DOCS = {}
LOCK = threading.Lock()

MAX_CONTEXT_CHARS = 12000


# -----------------------------
# Document extraction
# -----------------------------
def extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []

    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append((i, text))

    return pages


def extract_docx(path):
    from docx import Document

    doc = Document(str(path))
    lines = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            values = [value for value in values if value]
            if values:
                lines.append(" | ".join(values))

    return [(1, "\n".join(lines))]


def load_docs():
    with LOCK:
        DOCS.clear()

        for path in sorted(KNOWLEDGE.rglob("*")):
            if not path.is_file():
                continue

            suffix = path.suffix.lower()

            try:
                if suffix == ".pdf":
                    pages = extract_pdf(path)
                elif suffix == ".docx":
                    pages = extract_docx(path)
                else:
                    continue
            except Exception as exc:
                print(f"Warning: could not read {path.name}: {exc}")
                continue

            relative = path.relative_to(KNOWLEDGE)
            category = relative.parts[0] if len(relative.parts) > 1 else "Uncategorized"

            DOCS[str(relative)] = {
                "name": path.name,
                "category": category,
                "pages": pages,
            }


# -----------------------------
# Text chunking
# -----------------------------
def sentence_chunks(text, max_chars=900):
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []

    sentences = re.split(r"(?<=[.!?])\s+", text)

    chunks = []
    current = ""

    for sentence in sentences:
        if len(current) + len(sentence) + 1 <= max_chars:
            current = (current + " " + sentence).strip()
        else:
            if current:
                chunks.append(current)
            current = sentence

    if current:
        chunks.append(current)

    return chunks


# -----------------------------
# Category detection
# -----------------------------
CATEGORY_TERMS = {
    "03_Discipline": [
        "discipline", "disciplinary", "misconduct", "indiscipline",
        "punishment", "penalty", "hostel", "suspension", "expulsion",
        "appeal", "ragging", "conduct", "offence", "offense",
        "student misconduct", "disciplinary action",
    ],
    "02_Academic Rules": [
        "academic", "attendance", "registration", "semester", "course",
        "grade", "examination", "exam", "promotion", "degree",
        "undergraduate", "ug", "b.tech", "postgraduate", "pg",
        "m.tech", "m.sc", "ph.d", "phd", "research scholar",
    ],
    "04_Finance Procurement": [
        "procurement", "purchase", "tender", "gem", "gfr",
        "financial", "expenditure", "bid", "quotation", "sanction",
        "purchase committee", "payment",
    ],
    "01_Statutes_Gazette": [
        "statute", "statutes", "gazette", "act", "board of governors",
        "bog", "senate", "authority", "powers", "ordinance",
        "amendment", "statutory",
    ],
    "05_Office Order": [
        "office order", "order format", "office order format",
    ],
    "06_Noting": [
        "noting", "note", "file noting", "approval note", "proposal",
        "put up", "submitted for approval",
    ],
    "07_Minute": [
        "minutes", "minute", "mom", "meeting", "proceedings",
        "deliberation", "agenda",
    ],
    "08_Institutional Information": [
        "annual report", "institute", "department", "faculty",
        "staff", "campus", "institution", "vision", "mission",
    ],
}


def detect_categories(query):
    q = query.lower()
    scores = {}

    for category, words in CATEGORY_TERMS.items():
        score = 0
        for word in words:
            if word in q:
                score += 1
        if score:
            scores[category] = score

    return sorted(scores, key=lambda key: (-scores[key], key))


# -----------------------------
# Relevance search
# -----------------------------
def search_docs(query, limit=8):
    qlower = query.lower().strip()

    terms = [
        t.lower()
        for t in re.findall(r"[A-Za-z0-9]+", query)
        if len(t) >= 3
    ]

    preferred_categories = detect_categories(query)

    scored = []

    for doc_key, info in DOCS.items():
        name = info["name"]
        category = info["category"]
        pages = info["pages"]

        category_rank = (
            preferred_categories.index(category)
            if category in preferred_categories
            else 99
        )

        category_boost = 0
        if category in preferred_categories:
            category_boost = max(
                20 - (category_rank * 5),
                5,
            )

        name_lower = name.lower()

        for page_no, text in pages:
            if not text:
                continue

            for chunk in sentence_chunks(text):
                chunk_lower = chunk.lower()

                score = category_boost

                for term in terms:
                    occurrences = chunk_lower.count(term)
                    score += occurrences * 2

                    if term in name_lower:
                        score += 1

                if qlower and qlower in chunk_lower:
                    score += 25

                # Prefer chunks containing likely section words.
                if "process" in qlower and "process" in chunk_lower:
                    score += 8

                if "rule" in qlower and "rule" in chunk_lower:
                    score += 4

                if score > 0:
                    scored.append(
                        {
                            "score": score,
                            "document": name,
                            "category": category,
                            "page": page_no,
                            "snippet": chunk,
                        }
                    )

    scored.sort(
        key=lambda item: (
            -item["score"],
            0 if item["category"] in preferred_categories else 1,
            item["document"].lower(),
            item["page"],
        )
    )

    output = []
    seen = set()

    for hit in scored:
        key = (
            hit["document"],
            hit["page"],
            hit["snippet"][:180],
        )

        if key in seen:
            continue

        seen.add(key)
        output.append(hit)

        if len(output) >= limit:
            break

    return output


# -----------------------------
# Hermes integration
# -----------------------------
def call_hermes(prompt):
    """
    Calls the installed Hermes executable in one-shot mode.

    Uses the already configured Hermes provider/model, so no
    separate API key is required by this prototype.
    """

    try:
        result = subprocess.run(
            [
                "hermes",
                "-z",
                prompt,
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            cwd=str(BASE),
        )
    except FileNotFoundError:
        return (
            "Hermes could not be found in PATH. "
            "Open a new PowerShell window after installing Hermes "
            "and verify that `hermes --version` works."
        )
    except subprocess.TimeoutExpired:
        return "Hermes timed out while generating the response."
    except Exception as exc:
        return f"Hermes call failed: {exc}"

    if result.returncode != 0:
        details = (result.stderr or result.stdout or "").strip()
        return f"Hermes returned an error.\n\n{details}"

    return (result.stdout or "").strip()


def build_context(hits):
    parts = []

    for index, hit in enumerate(hits, start=1):
        parts.append(
            f"[SOURCE {index}]\n"
            f"Document: {hit['document']}\n"
            f"Category: {hit['category']}\n"
            f"Page: {hit['page']}\n"
            f"Passage: {hit['snippet']}"
        )

    context = "\n\n".join(parts)

    if len(context) > MAX_CONTEXT_CHARS:
        context = context[:MAX_CONTEXT_CHARS]

    return context


def answer_from_sources(query, hits):
    if not hits:
        return (
            "I could not locate a relevant passage in the local "
            "NIT Sikkim knowledge base."
        )

    context = build_context(hits)

    prompt = f"""
You are URI, the NIT Sikkim Administrative AI Assistant.

Answer the user's question using ONLY the supplied source passages.
Do not invent rules, clauses, authorities, dates, penalties, powers,
procedures, or institutional facts.

If the sources do not clearly establish something, say that it requires
verification or is not established by the supplied documents.

Distinguish carefully between:
- allegations and findings,
- committee recommendations and final decisions,
- suggested measures and mandatory provisions,
- examples/precedents and authoritative rules.

Give a concise, formal administrative answer.

At the end provide:
SOURCE(S): list the document name and page for the passages actually used.
VERIFICATION / APPROVAL POINTS: mention only genuine points requiring
human verification or competent-authority approval.

USER QUESTION:
{query}

SUPPLIED SOURCE PASSAGES:
{context}
""".strip()

    return call_hermes(prompt)


# -----------------------------
# Drafting helpers
# -----------------------------
def money_words(amount):
    try:
        number = int(round(float(str(amount).replace(",", ""))))
    except Exception:
        return ""

    ones = [
        "zero", "one", "two", "three", "four", "five", "six", "seven",
        "eight", "nine", "ten", "eleven", "twelve", "thirteen",
        "fourteen", "fifteen", "sixteen", "seventeen", "eighteen",
        "nineteen",
    ]

    tens = [
        "", "", "twenty", "thirty", "forty", "fifty",
        "sixty", "seventy", "eighty", "ninety",
    ]

    def under100(value):
        return (
            ones[value]
            if value < 20
            else tens[value // 10]
            + (" " + ones[value % 10] if value % 10 else "")
        )

    def under1000(value):
        if value >= 100:
            return (
                ones[value // 100]
                + " hundred"
                + (" " + under100(value % 100) if value % 100 else "")
            )
        return under100(value)

    if number < 1000:
        return under1000(number)

    if number < 100000:
        return (
            under1000(number // 1000)
            + " thousand"
            + (" " + under1000(number % 1000) if number % 1000 else "")
        )

    if number < 10000000:
        return (
            under1000(number // 100000)
            + " lakh"
            + (" " + under1000(number % 100000) if number % 100000 else "")
        )

    return str(number)


def draft_noting(data):
    subject = data.get("subject", "").strip() or "[Subject]"
    background = data.get("background", "").strip() or "[Background / reference]"
    requirement = (
        data.get("requirement", "").strip()
        or "[Requirement / proposal details]"
    )
    cost = data.get("cost", "").strip() or "[Amount, if applicable]"
    authority = (
        data.get("authority", "").strip()
        or "[Competent authority / approval level to be verified]"
    )
    extra = data.get("extra", "").strip()

    query = " ".join(
        [
            subject,
            background,
            requirement,
            "noting",
            "approval",
            "proposal",
        ]
    )

    hits = search_docs(query, 6)
    context = build_context(hits)

    prompt = f"""
You are URI, preparing a draft administrative noting for NIT Sikkim.

Use the source passages below only for institutional rules, procedures,
authority, and drafting conventions. Do not invent a rule or approval power.

Prepare a formal noting with:
1. Subject
2. Background / reference
3. Requirement / facts
4. Justification
5. Relevant rule/procedure, where supported
6. Financial implication, where applicable
7. Proposal
8. Approval requested
9. A short Verification / Approval Points section

Keep the distinction between factual information supplied by the officer
and conclusions derived from the sources.

INPUT:
Subject: {subject}
Background: {background}
Requirement: {requirement}
Estimated Cost: {cost}
Approval Level: {authority}
Additional Information: {extra}

SOURCE PASSAGES:
{context}
""".strip()

    answer = call_hermes(prompt)

    return answer, hits


def draft_office_order(instruction):
    instruction = instruction.strip() or "[Decision / direction to be issued]"

    hits = search_docs(
        instruction + " office order approval",
        6,
    )
    context = build_context(hits)

    prompt = f"""
You are URI, preparing a draft Office Order for NIT Sikkim.

Prepare a formal Office Order using the institutional style represented in
the source material.

Normally include:
- OFFICE ORDER
- clear operative decision/direction
- approval statement where appropriate
- appropriate signing designation
- Copy of Information to

Do not fabricate a reference number or date. Use placeholders where needed.

Do not convert a recommendation into a final decision unless the sources
establish approval/finality.

INPUT:
{instruction}

SOURCE PASSAGES:
{context}
""".strip()

    answer = call_hermes(prompt)

    return answer, hits


# -----------------------------
# HTML interface
# -----------------------------
PAGE = r"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>URI — NIT Sikkim Administrative AI Assistant</title>
<style>
body{
    font-family:Segoe UI,Arial,sans-serif;
    background:#f5f7fb;
    margin:0;
    color:#1f2937
}
header{
    background:#17324d;
    color:#fff;
    padding:22px 28px
}
header h1{margin:0;font-size:30px}
header p{margin:5px 0 0;opacity:.85}
.wrap{max-width:1150px;margin:24px auto;padding:0 18px}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:18px}
.card{
    background:#fff;
    border:1px solid #dbe2ea;
    border-radius:12px;
    padding:18px;
    box-shadow:0 2px 10px rgba(0,0,0,.04)
}
h2{margin-top:0}
.tabs button{
    border:0;
    background:#e9eef5;
    padding:10px 14px;
    border-radius:8px;
    margin:0 6px 6px 0;
    cursor:pointer
}
.tabs button.active{background:#17324d;color:#fff}
textarea,input{
    width:100%;
    box-sizing:border-box;
    padding:10px;
    border:1px solid #cbd5e1;
    border-radius:8px;
    margin:6px 0 12px;
    font:inherit
}
button.primary{
    background:#17324d;
    color:#fff;
    border:0;
    padding:11px 16px;
    border-radius:8px;
    cursor:pointer
}
button.primary:disabled{opacity:.6}
pre{
    white-space:pre-wrap;
    background:#f8fafc;
    border:1px solid #e2e8f0;
    padding:14px;
    border-radius:8px;
    max-height:650px;
    overflow:auto
}
.src{
    margin:8px 0;
    padding:10px;
    background:#f8fafc;
    border-left:4px solid #17324d
}
.muted{color:#64748b;font-size:13px}
.status{font-size:13px;margin-top:8px}
</style>
</head>
<body>
<header>
<h1>URI</h1>
<p>NIT Sikkim Administrative AI Assistant — local prototype</p>
</header>

<div class="wrap">
<div class="card">
<div class="tabs">
<button id="b1" class="active" onclick="show('noting',this)">Prepare Noting</button>
<button id="b2" onclick="show('rules',this)">Ask Rules</button>
<button id="b3" onclick="show('office',this)">Draft Office Order</button>
</div>

<div id="noting">
<h2>Prepare Noting</h2>
<input id="subject" placeholder="Subject">
<textarea id="background" rows="4" placeholder="Background / reference"></textarea>
<textarea id="requirement" rows="4" placeholder="Requirement / proposal details"></textarea>
<input id="cost" placeholder="Estimated cost, e.g. ₹8,00,000">
<input id="authority" placeholder="Approval level (optional)">
<textarea id="extra" rows="3" placeholder="Additional points (optional)"></textarea>
<button id="notingBtn" class="primary" onclick="draftNoting()">Generate Noting</button>
</div>

<div id="rules" style="display:none">
<h2>Ask Rules</h2>
<textarea id="query" rows="5" placeholder="Example: What does the Discipline Manual say about the disciplinary process?"></textarea>
<button id="rulesBtn" class="primary" onclick="askRules()">Ask Uri</button>
</div>

<div id="office" style="display:none">
<h2>Draft Office Order</h2>
<textarea id="oo" rows="8" placeholder="Describe the decision/direction to be issued."></textarea>
<button id="officeBtn" class="primary" onclick="draftOO()">Generate Office Order</button>
</div>
</div>

<div class="grid" style="margin-top:18px">
<div class="card">
<h2>Output</h2>
<pre id="output">Uri is ready.</pre>
<div id="status" class="status muted"></div>
</div>
<div class="card">
<h2>Sources</h2>
<div id="sources" class="muted">Sources used for rule-aware tasks will appear here.</div>
</div>
</div>
</div>

<script>
function show(id,button){
    for(const name of ['noting','rules','office']){
        document.getElementById(name).style.display=(name===id?'block':'none');
    }
    document.querySelectorAll('.tabs button').forEach(x=>x.classList.remove('active'));
    button.classList.add('active');
}

function setBusy(id,busy){
    const btn=document.getElementById(id);
    btn.disabled=busy;
    btn.textContent=busy?'Working...':btn.dataset.normal;
}

function initButtons(){
    document.querySelectorAll('button.primary').forEach(btn=>{
        btn.dataset.normal=btn.textContent;
    });
}
initButtons();

function renderSources(sources){
    const el=document.getElementById('sources');
    if(!sources || !sources.length){
        el.innerHTML='<span class="muted">No source passages located.</span>';
        return;
    }
    el.innerHTML=sources.map(s=>`
        <div class="src">
            <strong>${escapeHtml(s.document)}</strong><br>
            ${escapeHtml('Page ' + s.page)}<br>
            <span class="muted">${escapeHtml(s.snippet)}</span>
        </div>`).join('');
}

function escapeHtml(value){
    return String(value)
      .replaceAll('&','&amp;')
      .replaceAll('<','&lt;')
      .replaceAll('>','&gt;')
      .replaceAll('"','&quot;')
      .replaceAll("'","&#039;");
}

async function askRules(){
    const query=document.getElementById('query').value.trim();
    if(!query) return;

    setBusy('rulesBtn',true);
    document.getElementById('status').textContent='Searching local knowledge base and asking Hermes...';

    try{
        const r=await fetch('/api/rules',{
            method:'POST',
            headers:{'Content-Type':'application/json'},
            body:JSON.stringify({query})
        });
        const data=await r.json();

        document.getElementById('output').textContent=data.answer || data.error || '';
        renderSources(data.sources || []);
        document.getElementById('status').textContent='Completed.';
    }catch(e){
        document.getElementById('output').textContent='Error: '+e;
        document.getElementById('status').textContent='';
    }finally{
        setBusy('rulesBtn',false);
    }
}

async function draftNoting(){
    const payload={
        subject:document.getElementById('subject').value,
        background:document.getElementById('background').value,
        requirement:document.getElementById('requirement').value,
        cost:document.getElementById('cost').value,
        authority:document.getElementById('authority').value,
        extra:document.getElementById('extra').value
    };

    setBusy('notingBtn',true);
    document.getElementById('status').textContent='Preparing noting with Hermes...';

    try{
        const r=await fetch('/api/noting',{
            method:'POST',
            headers:{'Content-Type':'application/json'},
            body:JSON.stringify(payload)
        });
        const data=await r.json();

        document.getElementById('output').textContent=data.answer || data.error || '';
        renderSources(data.sources || []);
        document.getElementById('status').textContent='Completed.';
    }catch(e){
        document.getElementById('output').textContent='Error: '+e;
        document.getElementById('status').textContent='';
    }finally{
        setBusy('notingBtn',false);
    }
}

async function draftOO(){
    const instruction=document.getElementById('oo').value.trim();
    if(!instruction) return;

    setBusy('officeBtn',true);
    document.getElementById('status').textContent='Preparing Office Order with Hermes...';

    try{
        const r=await fetch('/api/office-order',{
            method:'POST',
            headers:{'Content-Type':'application/json'},
            body:JSON.stringify({instruction})
        });
        const data=await r.json();

        document.getElementById('output').textContent=data.answer || data.error || '';
        renderSources(data.sources || []);
        document.getElementById('status').textContent='Completed.';
    }catch(e){
        document.getElementById('output').textContent='Error: '+e;
        document.getElementById('status').textContent='';
    }finally{
        setBusy('officeBtn',false);
    }
}
</script>
</body>
</html>
"""


# -----------------------------
# HTTP server
# -----------------------------
class Handler(BaseHTTPRequestHandler):
    def _send(self, code, body, content_type="text/html; charset=utf-8"):
        encoded = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)

    def _json_body(self):
        length = int(self.headers.get("Content-Length", "0"))
        raw = self.rfile.read(length)
        return json.loads(raw.decode("utf-8"))

    def do_GET(self):
        path = urlparse(self.path).path

        if path == "/":
            self._send(200, PAGE)
            return

        if path == "/api/health":
            self._send(
                200,
                json.dumps(
                    {
                        "ok": True,
                        "documents": len(DOCS),
                    }
                ),
                "application/json; charset=utf-8",
            )
            return

        self._send(404, "Not found", "text/plain; charset=utf-8")

    def do_POST(self):
        path = urlparse(self.path).path

        try:
            data = self._json_body()

            if path == "/api/rules":
                query = data.get("query", "").strip()
                hits = search_docs(query, 8)
                answer = answer_from_sources(query, hits)

                self._send(
                    200,
                    json.dumps(
                        {
                            "answer": answer,
                            "sources": hits,
                        },
                    ),
                    "application/json; charset=utf-8",
                )
                return

            if path == "/api/noting":
                answer, hits = draft_noting(data)

                self._send(
                    200,
                    json.dumps(
                        {
                            "answer": answer,
                            "sources": hits,
                        },
                    ),
                    "application/json; charset=utf-8",
                )
                return

            if path == "/api/office-order":
                instruction = data.get("instruction", "").strip()
                answer, hits = draft_office_order(instruction)

                self._send(
                    200,
                    json.dumps(
                        {
                            "answer": answer,
                            "sources": hits,
                        },
                    ),
                    "application/json; charset=utf-8",
                )
                return

            self._send(404, "Not found", "text/plain; charset=utf-8")

        except Exception as exc:
            self._send(
                500,
                json.dumps({"error": str(exc)}),
                "application/json; charset=utf-8",
            )


if __name__ == "__main__":
    print("Loading NIT Sikkim knowledge base...")
    load_docs()
    print(f"Loaded {len(DOCS)} documents.")
    print("URI V2 is ready.")
    print("Open http://127.0.0.1:8000 in your browser.")

    server = ThreadingHTTPServer(("127.0.0.1", 8000), Handler)
    server.serve_forever()
