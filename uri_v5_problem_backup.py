import json, re, subprocess, threading
from datetime import date
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse

BASE=Path(__file__).resolve().parent
KNOWLEDGE=BASE/'knowledge'
LIBRARY=BASE/'institutional_library'
MEMORY=BASE/'uri_memory.json'
DOCS={}; LOCK=threading.Lock(); MAX_CONTEXT=16000

CAT_TERMS={
'03_Discipline':['discipline','disciplinary','misconduct','indiscipline','punishment','penalty','suspension','expulsion','appeal','ragging'],
'02_Academic Rules':['academic','attendance','registration','semester','course','grade','examination','exam','promotion','degree','ug','b.tech','pg','m.tech','m.sc','ph.d','phd'],
'04_Finance Procurement':['procurement','purchase','tender','gem','gfr','financial','expenditure','bid','quotation','sanction','payment','delegation'],
'01_Statutes_Gazette':['statute','statutes','gazette','act','board of governors','bog','senate','authority','powers','ordinance','amendment'],
'09_Government_Office_Procedure':['csmop','office procedure','file management','noting guidelines','drafting']}
LIB_TERMS={
'01_Noting':['noting','note','file noting','approval note'],
'02_Office Orders':['office order','order'],
'03_Notices':['notice'],
'04_MoM':['minutes','minute','mom','meeting'],
'05_Minesterial Replies':['ministry','ministerial','lok sabha','rajya sabha','parliamentary','question'],
'06_Letters':['letter','correspondence'],
'07_Action Taken Report':['action taken report','atr','action taken'],
'08_others':['draft','format','precedent','agreement','certificate']}

def extract_pdf(p):
 from pypdf import PdfReader
 r=PdfReader(str(p)); return [(i,x.extract_text() or '') for i,x in enumerate(r.pages,1)]

def extract_docx(p):
 from docx import Document
 d=Document(str(p)); lines=[]
 for q in d.paragraphs:
  if q.text.strip(): lines.append(q.text.strip())
 for t in d.tables:
  for row in t.rows:
   vals=[c.text.strip() for c in row.cells if c.text.strip()]
   if vals: lines.append(' | '.join(vals))
 return [(1,'\n'.join(lines))]

def add_root(root, source):
 if not root.exists(): return
 for p in sorted(root.rglob('*')):
  if not p.is_file() or p.suffix.lower() not in {'.pdf','.docx'}: continue
  try: pages=extract_pdf(p) if p.suffix.lower()=='.pdf' else extract_docx(p)
  except Exception as e: print(f'Warning: could not read {p.name}: {e}'); continue
  rel=p.relative_to(root); cat=rel.parts[0] if len(rel.parts)>1 else 'Uncategorized'
  text='\n'.join(t for _,t in pages); years=re.findall(r'\b(?:19|20)\d{2}\b',f'{p.name} {text[:5000]}')
  DOCS[f'{source}:{rel}']={'name':p.name,'relative':str(rel),'category':cat,'source_type':source,'role':('AUTHORITATIVE / REFERENCE' if source=='authoritative' else 'INSTITUTIONAL PRECEDENT / DRAFTING EXAMPLE'),'year':years[-1] if years else '','pages':pages}

def load_docs():
 with LOCK:
  DOCS.clear(); add_root(KNOWLEDGE,'authoritative'); add_root(LIBRARY,'institutional_library')

def chunks(text, size=950):
 text=re.sub(r'\s+',' ',text).strip()
 if not text:return []
 ss=re.split(r'(?<=[.!?])\s+',text); out=[]; cur=''
 for s in ss:
  if len(cur)+len(s)+1<=size: cur=(cur+' '+s).strip()
  else:
   if cur: out.append(cur)
   cur=s
 if cur: out.append(cur)
 return out

def prefs(q, mapping):
 q=q.lower(); scores={c:sum(1 for w in ws if w in q) for c,ws in mapping.items()}
 return sorted([c for c,n in scores.items() if n],key=lambda c:(-scores[c],c))

def search(q, limit=10, precedent=False):
 q=q.lower().strip(); terms=[t for t in re.findall(r'[a-z0-9]+',q) if len(t)>=3]
 ap=prefs(q,CAT_TERMS); lp=prefs(q,LIB_TERMS)
 rule=any(x in q for x in ['what rule','which rule','applicable rule','regulation','provision','discipline manual','statute','gfr','competent authority','delegation','procedure'])
 ranked=[]
 for info in DOCS.values():
  cat=info['category']; src=info['source_type']; name=info['name'].lower()
  for page,text in info['pages']:
   for ch in chunks(text):
    cl=ch.lower(); score=sum(cl.count(t)*2+(2 if t in name else 0) for t in terms)
    if cat in ap: score+=max(35-ap.index(cat)*8,8)
    if cat in lp: score+=max(28-lp.index(cat)*5,5)
    if q and q in cl: score+=35
    if rule: score+=30 if src=='authoritative' else -8
    if precedent: score+=36 if src=='institutional_library' else 3
    if score>0: ranked.append({'score':score,'document':info['name'],'relative':info['relative'],'category':cat,'year':info['year'],'source_type':src,'role':info['role'],'page':page,'snippet':ch})
 ranked.sort(key=lambda x:(-x['score'],x['document'].lower(),x['page']))
 out=[]; seen=set()
 for h in ranked:
  k=(h['document'],h['page'],h['snippet'][:180])
  if k in seen: continue
  seen.add(k); out.append(h)
  if len(out)>=limit: break
 return out

def context(hits):
 return '\n\n'.join(f"[SOURCE {i}]\nDocument: {h['document']}\nCategory: {h['category']}\nYear: {h['year'] or 'Not identified'}\nRole: {h['role']}\nPage: {h['page']}\nPassage: {h['snippet']}" for i,h in enumerate(hits,1))[:MAX_CONTEXT]

def hermes(prompt):
 try:
  r=subprocess.run(['hermes','-z',prompt],capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=120,cwd=str(BASE))
 except FileNotFoundError: return 'Hermes was not found in PATH. Confirm `hermes --version` works in a new PowerShell.'
 except subprocess.TimeoutExpired: return 'Hermes timed out. Please try again.'
 except Exception as e: return f'Hermes call failed: {e}'
 return (r.stdout if r.returncode==0 else 'Hermes returned an error.\n\n'+(r.stderr or r.stdout)).strip()

def load_memory():
 if not MEMORY.exists(): return []
 try:
  x=json.loads(MEMORY.read_text(encoding='utf-8')); return x if isinstance(x,list) else []
 except Exception: return []

def memory_text():
 m=load_memory()
 return '\n'.join(f"- {x.get('fact')} | status={x.get('status')} | confirmed={x.get('confirmed')}" for x in m) if m else '(none)'

def chat(message,history):
 previous=any(x in message.lower() for x in ['previous','past','earlier','historical','precedent','similar','find old'])
 hits=search(message,10,previous)
 hist='\n'.join(f"{m.get('role','user').upper()}: {m.get('content','')}" for m in history[-10:])
 prompt=f'''You are URI, the NIT Sikkim Administrative AI Assistant.\n\nUsers can describe administrative goals naturally. Decide whether to answer, search, draft, or ask for missing information.\n\nRULES: Never invent facts, rules, authorities, dates, penalties, procedures, approvals, designations or recipients. Distinguish current authoritative sources from historical NIT Sikkim drafting examples. Historical examples show style/past practice and do not automatically establish current authority. Do not infer competent authority from amount or designation. Distinguish recommendations, approvals and final decisions. Use memory only as convenience; ask for confirmation if time-sensitive information may be stale. Ask only essential questions for drafting.\n\nAPPROVED MEMORY:\n{memory_text()}\n\nRECENT CONVERSATION:\n{hist or '(none)'}\n\nUSER:\n{message}\n\nLOCAL SOURCES:\n{context(hits) if hits else '(none)'}\n\nRespond naturally and concisely. For drafting requests, ask essential missing questions before finalizing. For rules questions, answer from sources. For historical-example requests, identify the most relevant previous documents. If a new stable, non-sensitive, clearly confirmed institutional fact should be remembered, append exactly: MEMORY SUGGESTION: [fact].'''
 return hermes(prompt),hits

PAGE='''<!doctype html><html><head><meta charset="utf-8"><title>URI</title><style>body{margin:0;background:#f5f7fb;color:#1f2937;font-family:Segoe UI,Arial,sans-serif}header{background:#17324d;color:#fff;padding:20px 24px}header h1{margin:0;font-size:28px}header p{margin:4px 0 0;opacity:.85}.wrap{max-width:1050px;margin:auto;padding:18px}.chat{height:70vh;overflow:auto;padding:5px}.msg{display:flex;margin:12px 0}.bubble{max-width:82%;padding:12px 15px;border-radius:14px;white-space:pre-wrap;line-height:1.45}.user{justify-content:flex-end}.user .bubble{background:#17324d;color:#fff}.assistant .bubble{background:#fff;border:1px solid #dbe2ea}.composer{display:flex;gap:10px;background:#fff;border:1px solid #dbe2ea;padding:10px;border-radius:14px}.composer textarea{flex:1;resize:none;border:0;outline:0;font:inherit;min-height:52px}.send{border:0;background:#17324d;color:#fff;padding:0 20px;border-radius:10px;cursor:pointer}.send:disabled{opacity:.6}.tools{display:flex;gap:8px;margin-bottom:8px}.tool{border:1px solid #cbd5e1;background:#fff;border-radius:8px;padding:7px 10px;cursor:pointer}.src{background:#fff;border:1px solid #e2e8f0;border-radius:8px;padding:9px;margin:7px 0;font-size:12px}.status{font-size:12px;color:#64748b;margin:8px 4px}</style></head><body><header><h1>URI</h1><p>NIT Sikkim Administrative AI Assistant</p></header><div class="wrap"><div class="tools"><button class="tool" onclick="newChat()">New chat</button><button class="tool" onclick="mem()">View approved memory</button></div><div id="chat" class="chat"></div><div id="status" class="status">Ready.</div><div class="composer"><textarea id="input" placeholder="Tell Uri what you need..."></textarea><button id="send" class="send" onclick="send()">Send</button></div><div id="sources"></div></div><script>let history=[];const box=document.getElementById('chat');function esc(v){return String(v).replaceAll('&','&amp;').replaceAll('<','&lt;').replaceAll('>','&gt;')}function add(role,text){const r=document.createElement('div');r.className='msg '+role;const b=document.createElement('div');b.className='bubble';b.innerHTML=esc(text);r.appendChild(b);box.appendChild(r);box.scrollTop=box.scrollHeight}function newChat(){history=[];box.innerHTML='';document.getElementById('sources').innerHTML='';add('assistant','Hello. I am Uri. Tell me what you need.')}async function send(){const i=document.getElementById('input'),text=i.value.trim();if(!text)return;add('user',text);history.push({role:'user',content:text});i.value='';document.getElementById('send').disabled=true;document.getElementById('status').textContent='Uri is working...';try{const r=await fetch('/api/chat',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({message:text,history})});const d=await r.json();add('assistant',d.answer||d.error||'No response.');history.push({role:'assistant',content:d.answer||''});render(d.sources||[])}catch(e){add('assistant','Connection error: '+e)}finally{document.getElementById('send').disabled=false;document.getElementById('status').textContent='Ready.'}}function render(s){const e=document.getElementById('sources');if(!s.length){e.innerHTML='';return}e.innerHTML='<h3>Sources</h3>'+s.map(x=>'<div class="src"><b>'+esc(x.document)+'</b> - page '+esc(x.page)+'<br>'+esc(x.role)+'<br>'+esc(x.snippet)+'</div>').join('')}async function mem(){try{const d=await (await fetch('/api/memory')).json();add('assistant',d.memory.length?d.memory.map(x=>x.fact+' ['+x.status+', confirmed '+x.confirmed+']').join('\n'):'No approved memory stored.')}catch(e){add('assistant','Could not read memory.')}}document.getElementById('input').addEventListener('keydown',e=>{if(e.key==='Enter'&&!e.shiftKey){e.preventDefault();send()}});newChat();</script></body></html>'''

class Handler(BaseHTTPRequestHandler):
 def send_data(self,code,body,ctype='text/html; charset=utf-8'):
  data=body.encode('utf-8'); self.send_response(code); self.send_header('Content-Type',ctype); self.send_header('Content-Length',str(len(data))); self.end_headers(); self.wfile.write(data)
 def body(self):
  n=int(self.headers.get('Content-Length','0')); return json.loads(self.rfile.read(n).decode('utf-8'))
 def do_GET(self):
  p=urlparse(self.path).path
  if p=='/': self.send_data(200,PAGE); return
  if p=='/api/health':
   a=sum(x['source_type']=='authoritative' for x in DOCS.values()); i=sum(x['source_type']=='institutional_library' for x in DOCS.values())
   self.send_data(200,json.dumps({'ok':True,'total':len(DOCS),'authoritative':a,'institutional':i,'memory':len(load_memory())}),'application/json'); return
  if p=='/api/memory': self.send_data(200,json.dumps({'memory':load_memory()},ensure_ascii=False),'application/json'); return
  self.send_data(404,'Not found','text/plain')
 def do_POST(self):
  p=urlparse(self.path).path
  try:
   d=self.body()
   if p=='/api/chat':
    ans,hits=chat(d.get('message','').strip(),d.get('history',[])); self.send_data(200,json.dumps({'answer':ans,'sources':hits},ensure_ascii=False),'application/json'); return
   if p=='/api/memory/confirm':
    fact=d.get('fact','').strip()
    if not fact: self.send_data(400,json.dumps({'error':'fact required'}),'application/json'); return
    m=load_memory(); m.append({'fact':fact,'status':'CONFIRMED','confirmed':str(date.today())}); MEMORY.write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8'); self.send_data(200,json.dumps({'ok':True}),'application/json'); return
   self.send_data(404,'Not found','text/plain')
  except Exception as e: self.send_data(500,json.dumps({'error':str(e)}),'application/json')

if __name__=='__main__':
 print('Loading NIT Sikkim knowledge base and institutional library...'); load_docs()
 a=sum(x['source_type']=='authoritative' for x in DOCS.values()); i=sum(x['source_type']=='institutional_library' for x in DOCS.values())
 print(f'Loaded {len(DOCS)} total documents.'); print(f'Authoritative/reference documents: {a}'); print(f'Institutional-library documents: {i}'); print('URI V5 - Conversational Agent'); print('Open http://127.0.0.1:8000 in your browser.')
 ThreadingHTTPServer(('127.0.0.1',8000),Handler).serve_forever()
