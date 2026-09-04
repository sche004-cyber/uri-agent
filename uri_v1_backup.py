import os, re, json, html, threading
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse

BASE = Path(__file__).resolve().parent
KNOWLEDGE = BASE / 'knowledge'
CACHE = BASE / 'cache'
CACHE.mkdir(exist_ok=True)

DOCS = {}
LOCK = threading.Lock()


def extract_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages=[]
    for i,p in enumerate(reader.pages, start=1):
        try: t=p.extract_text() or ''
        except Exception: t=''
        pages.append((i,t))
    return pages


def extract_docx(path):
    from docx import Document
    doc=Document(str(path))
    lines=[]
    for p in doc.paragraphs:
        txt=p.text.strip()
        if txt: lines.append(txt)
    for table in doc.tables:
        for row in table.rows:
            vals=[c.text.strip() for c in row.cells]
            vals=[v for v in vals if v]
            if vals: lines.append(' | '.join(vals))
    return [(1,'\n'.join(lines))]


def load_docs():
    with LOCK:
        DOCS.clear()

        for path in sorted(KNOWLEDGE.rglob('*')):
            if not path.is_file():
                continue

            suffix = path.suffix.lower()

            try:
                if suffix == '.pdf':
                    pages = extract_pdf(path)
                elif suffix == '.docx':
                    pages = extract_docx(path)
                else:
                    continue
            except Exception as e:
                print(f"Warning: could not read {path.name}: {e}")
                continue

            relative = path.relative_to(KNOWLEDGE)
            category = relative.parts[0] if len(relative.parts) > 1 else "Uncategorized"

            DOCS[str(relative)] = {
                "name": path.name,
                "category": category,
                "pages": pages
            }


def sentence_chunks(text, max_chars=700):
    text=re.sub(r'\s+',' ',text).strip()
    if not text: return []
    sents=re.split(r'(?<=[.!?])\s+', text)
    out=[]; cur=''
    for s in sents:
        if len(cur)+len(s)+1 <= max_chars: cur=(cur+' '+s).strip()
        else:
            if cur: out.append(cur)
            cur=s
    if cur: out.append(cur)
    return out


def search_docs(query, limit=8):
    qlower = query.lower().strip()

    terms = [
        t.lower()
        for t in re.findall(r'[A-Za-z0-9]+', query)
        if len(t) >= 3
    ]

    # Detect the likely subject area.
    category_terms = {
        "03_Discipline": [
            "discipline", "misconduct", "indiscipline", "punishment",
            "penalty", "student conduct", "hostel", "suspension",
            "expulsion", "appeal", "disciplinary"
        ],
        "02_Academic Rules": [
            "academic", "attendance", "registration", "semester",
            "course", "grade", "examination", "promotion", "degree",
            "ug", "b.tech", "pg", "m.tech", "m.sc", "ph.d", "phd"
        ],
        "04_Finance Procurement": [
            "procurement", "purchase", "tender", "gem", "gfr",
            "financial", "expenditure", "bid", "quotation"
        ],
        "01_Statutes_Gazette": [
            "statute", "statutes", "gazette", "act", "board of governors",
            "bog", "senate", "authority", "powers"
        ],
        "05_Office Order": [
            "office order", "order format", "office order format"
        ],
        "06_Noting": [
            "noting", "note", "file noting", "approval note", "proposal"
        ],
        "07_Minute": [
            "minutes", "minute", "mom", "meeting", "proceedings"
        ],
        "08_Institutional Information": [
            "annual report", "institute", "department", "faculty",
            "staff", "campus", "institution"
        ],
    }

    category_scores = {}
    for category, words in category_terms.items():
        category_scores[category] = sum(3 for w in words if w in qlower)

    strongest_category = None
    strongest_score = 0

    for category, score in category_scores.items():
        if score > strongest_score:
            strongest_category = category
            strongest_score = score

    scored = []

    for doc_key, info in DOCS.items():
        name = info["name"]
        category = info["category"]
        pages = info["pages"]

        for page_no, text in pages:
            if not text:
                continue

            chunks = sentence_chunks(text)

            for ch in chunks:
                cl = ch.lower()

                score = 0

                # Ordinary lexical relevance.
                for term in terms:
                    count = cl.count(term)
                    score += count * 2

                # Exact query phrase.
                if qlower and qlower in cl:
                    score += 12

                # Strong boost when the query's subject matches the
                # document category.
                if strongest_category and category == strongest_category:
                    score += strongest_score * 4

                # Additional filename relevance.
                for term in terms:
                    if term in name.lower():
                        score += 2

                if score > 0:
                    scored.append(
                        (score, name, category, page_no, ch)
                    )

    scored.sort(
        key=lambda x: (-x[0], x[1].lower(), x[3])
    )

    out = []
    seen = set()

    for score, name, category, page, ch in scored:
        key = (name, page, ch[:150])

        if key in seen:
            continue

        seen.add(key)

        out.append({
            "score": score,
            "document": name,
            "category": category,
            "page": page,
            "snippet": ch
        })

        if len(out) >= limit:
            break

    return out


def money_words(amount):
    try:
        n=int(round(float(str(amount).replace(',',''))))
    except Exception:
        return ''
    ones=['zero','one','two','three','four','five','six','seven','eight','nine','ten','eleven','twelve','thirteen','fourteen','fifteen','sixteen','seventeen','eighteen','nineteen']
    tens=['','','twenty','thirty','forty','fifty','sixty','seventy','eighty','ninety']
    def under100(x): return ones[x] if x<20 else tens[x//10]+(' '+ones[x%10] if x%10 else '')
    def under1000(x): return (ones[x//100]+' hundred'+(' '+under100(x%100) if x%100 else '')) if x>=100 else under100(x)
    if n<1000: return under1000(n)
    if n<100000: return under1000(n//1000)+' thousand'+(' '+under1000(n%1000) if n%1000 else '')
    if n<10000000: return under1000(n//100000)+' lakh'+(' '+under1000(n%100000) if n%100000 else '')
    return str(n)


def draft_noting(data):
    subject=data.get('subject','').strip() or '[Subject]'
    background=data.get('background','').strip() or '[Background / reference to be inserted]'
    requirement=data.get('requirement','').strip() or '[Requirement / proposal details]'
    cost=data.get('cost','').strip() or '[Amount, if applicable]'
    authority=data.get('authority','').strip() or '[Competent authority / approval level to be verified]'
    extra=data.get('extra','').strip()
    q=' '.join([subject, requirement, 'procurement', 'approval'])
    hits=search_docs(q,6)
    sources='\n'.join([f"- {h['document']} (page {h['page']}): {h['snippet']}" for h in hits[:5]]) or '- No relevant source located in the local knowledge base.'
    rupees=''
    if re.search(r'\d',cost):
        nums=re.sub(r'[^0-9.]','',cost)
        try:
            rupees=f" (Rupees {money_words(nums)} only)"
        except Exception: pass
    draft=f'''NOTE\n\nSubject: {subject}\n\n1. Background\n{background}\n\n2. Requirement / Proposal\n{requirement}\n\n3. Financial Implication\nThe estimated financial implication is {cost}{rupees}.\n\n4. Rules / Procedure\nThe applicable institutional rules and Government instructions should be verified against the subject matter before approval. Uri located the following potentially relevant material:\n{sources}\n\n5. Proposal\nIn view of the above, approval is solicited for {subject.lower()}. The appropriate procurement / administrative procedure and competent authority may be confirmed before further action.\n\n{extra + chr(10) if extra else ''}Submitted for approval.\n\nApproval required: {authority}\n'''
    return draft, hits


HTML = r'''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Uri - NIT Sikkim Office AI</title><style>
body{font-family:Segoe UI,Arial,sans-serif;background:#f5f7fb;margin:0;color:#1f2937}header{background:#17324d;color:#fff;padding:22px 28px}header h1{margin:0;font-size:30px}header p{margin:5px 0 0;opacity:.85}.wrap{max-width:1100px;margin:24px auto;padding:0 18px}.grid{display:grid;grid-template-columns:1fr 1fr;gap:18px}.card{background:white;border:1px solid #dbe2ea;border-radius:12px;padding:18px;box-shadow:0 2px 10px rgba(0,0,0,.04)}h2{margin-top:0}.tabs button{border:0;background:#e9eef5;padding:10px 14px;border-radius:8px;margin:0 6px 6px 0;cursor:pointer}.tabs button.active{background:#17324d;color:#fff}textarea,input{width:100%;box-sizing:border-box;padding:10px;border:1px solid #cbd5e1;border-radius:8px;margin:6px 0 12px;font:inherit}button.primary{background:#17324d;color:#fff;border:0;padding:11px 16px;border-radius:8px;cursor:pointer}pre{white-space:pre-wrap;background:#f8fafc;border:1px solid #e2e8f0;padding:14px;border-radius:8px;max-height:520px;overflow:auto}.src{margin:8px 0;padding:10px;background:#f8fafc;border-left:4px solid #17324d}.muted{color:#64748b;font-size:13px}.status{font-size:13px;margin-top:8px}</style></head><body><header><h1>URI</h1><p>NIT Sikkim Administrative AI Assistant — local prototype</p></header><div class="wrap"><div class="card"><div class="tabs"><button id="b1" class="active" onclick="show('noting',this)">Prepare Noting</button><button id="b2" onclick="show('rules',this)">Ask Rules</button><button id="b3" onclick="show('office',this)">Draft Office Order</button></div><div id="noting"><h2>Prepare Noting</h2><input id="subject" placeholder="Subject"><textarea id="background" rows="4" placeholder="Background / reference"></textarea><textarea id="requirement" rows="4" placeholder="Requirement / proposal details"></textarea><input id="cost" placeholder="Estimated cost, e.g. ₹8,00,000"><input id="authority" placeholder="Approval level (optional)"><textarea id="extra" rows="3" placeholder="Additional points (optional)"></textarea><button class="primary" onclick="draftNoting()">Generate Noting</button></div><div id="rules" style="display:none"><h2>Ask Rules</h2><textarea id="query" rows="5" placeholder="Example: What does the Discipline Manual say about the disciplinary process?"></textarea><button class="primary" onclick="askRules()">Search Knowledge Base</button></div><div id="office" style="display:none"><h2>Draft Office Order</h2><textarea id="oo" rows="8" placeholder="Describe the decision/direction to be issued."></textarea><button class="primary" onclick="draftOO()">Generate Office Order</button></div></div><div class="grid" style="margin-top:18px"><div class="card"><h2>Output</h2><pre id="output">Uri is ready.</pre><div id="status" class="status muted"></div></div><div class="card"><h2>Sources</h2><div id="sources" class="muted">Sources used for rule-aware tasks will appear here.</div></div></div></div><script>
function show(id,btn){['noting','rules','office'].forEach(x=>document.getElementById(x).style.display=x===id?'block':'none');document.querySelectorAll('.tabs button').forEach(b=>b.classList.remove('active'));btn.classList.add('active')}
async function post(path,data){let r=await fetch(path,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(data)});return await r.json()}
function renderSources(h){if(!h||!h.length){document.getElementById('sources').innerHTML='No local sources found.';return}document.getElementById('sources').innerHTML=h.map(x=>`<div class="src"><b>${esc(x.document)}</b> — page ${x.page}<br>${esc(x.snippet)}</div>`).join('')}
function esc(s){return s.replace(/[&<>\"]/g,m=>({'&':'&amp;','<':'&lt;','>':'&gt;','\"':'&quot;'}[m]))}
async function draftNoting(){let d=await post('/api/draft-noting',{subject:subject.value,background:background.value,requirement:requirement.value,cost:cost.value,authority:authority.value,extra:extra.value});output.textContent=d.draft;renderSources(d.sources);status.textContent='Draft mode: verify applicable rules and approval authority before issuing.'}
async function askRules(){let d=await post('/api/search',{query:query.value});output.textContent=d.summary||'Search results';renderSources(d.results);status.textContent='Search is local and source-based; absence of a result does not prove absence of a rule.'}
async function draftOO(){let d=await post('/api/draft-office-order',{instruction:oo.value});output.textContent=d.draft;renderSources(d.sources);status.textContent='Office Order is a draft only and requires competent-authority review.'}
</script></body></html>'''

class Handler(BaseHTTPRequestHandler):
    def _send(self, code, body, ctype='text/html; charset=utf-8'):
        self.send_response(code); self.send_header('Content-Type',ctype); self.send_header('Content-Length',str(len(body))); self.end_headers(); self.wfile.write(body)
    def do_GET(self):
        p=urlparse(self.path).path
        if p in ('/','/index.html'): return self._send(200,HTML.encode())
        if p=='/api/health': return self._send(200,json.dumps({'ok':True,'documents':len(DOCS)}).encode(),'application/json')
        self._send(404,b'Not found')
    def do_POST(self):
        length=int(self.headers.get('Content-Length','0')); raw=self.rfile.read(length)
        try: data=json.loads(raw or b'{}')
        except Exception: return self._send(400,b'{"error":"Invalid JSON"}','application/json')
        p=urlparse(self.path).path
        if p=='/api/search':
            q=data.get('query','').strip(); results=search_docs(q,8)
            summary='No matching source located in the local knowledge base.' if not results else 'Top matching passages from the local NIT Sikkim knowledge base.'
            return self._send(200,json.dumps({'summary':summary,'results':results},ensure_ascii=False).encode(),'application/json')
        if p=='/api/draft-noting':
            draft,hits=draft_noting(data); return self._send(200,json.dumps({'draft':draft,'sources':hits},ensure_ascii=False).encode(),'application/json')
        if p=='/api/draft-office-order':
            instruction=data.get('instruction','').strip() or '[Direction / decision]'
            hits=search_docs(instruction+' office order approval',6)
            draft=f'''NATIONAL INSTITUTE OF TECHNOLOGY SIKKIM\n\nRef: [To be assigned]\nDate: [Date]\n\nOFFICE ORDER\n\n{instruction}\n\nThis is issued with the approval of the competent authority.\n\n[Authorized Signatory]\n[Designation]\n\nCopy of Information to:\n1. Director's Office\n2. Concerned officer(s)/section(s)\n3. File\n'''
            return self._send(200,json.dumps({'draft':draft,'sources':hits},ensure_ascii=False).encode(),'application/json')
        self._send(404,b'Not found')

if __name__=='__main__':
    print('Loading NIT Sikkim knowledge base...')
    load_docs()
    print(f'Loaded {len(DOCS)} documents.')
    print('Open http://127.0.0.1:8000 in your browser.')
    ThreadingHTTPServer(('127.0.0.1',8000),Handler).serve_forever()
