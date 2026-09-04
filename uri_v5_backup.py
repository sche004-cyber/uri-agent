import json
import re
import subprocess
import threading
from datetime import date
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse

# ============================================================
# URI V5 - NIT Sikkim Administrative AI Assistant
#
# One chat interface
# Two knowledge sources:
#   1. knowledge/              = authoritative/reference material
#   2. institutional_library/  = historical NIT Sikkim examples
#
# Hermes is used as the AI reasoning engine.
# ============================================================

BASE = Path(__file__).resolve().parent
KNOWLEDGE = BASE / "knowledge"
INSTITUTIONAL_LIBRARY = BASE / "institutional_library"
MEMORY_FILE = BASE / "uri_memory.json"

DOCS = {}
LOCK = threading.Lock()

MAX_CONTEXT_CHARS = 16000
MAX_HISTORY_MESSAGES = 12


# ------------------------------------------------------------
# Document extraction
# ------------------------------------------------------------
def extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        pages.append((page_number, text))

    return pages


def extract_docx(path):
    from docx import Document

    doc = Document(str(path))
    lines = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            values = [value for value in values if value]
            if values:
                lines.append(" | ".join(values))

    return [(1, "\n".join(lines))]


def infer_year(text, filename):
    years = re.findall(
        r"\b(?:19|20)\d{2}\b",
        f"{filename} {text[:8000]}"
    )

    if not years:
        return ""

    return max(years)


def infer_document_type(category, filename):
    value = f"{category} {filename}".lower()

    checks = [
        ("Office Order", ["office order"]),
        ("Noting", ["noting", "file noting", "approval note"]),
        ("Minutes of Meeting", ["minutes", "minute", "mom"]),
        ("Notice", ["notice"]),
        ("Letter", ["letter", "correspondence"]),
        ("Statutory / Gazette", ["statute", "gazette", "act"]),
        ("Rules / Manual", ["rule", "regulation", "manual"]),
        ("Finance / Procurement", ["gfr", "finance", "procurement"]),
    ]

    for document_type, words in checks:
        if any(word in value for word in words):
            return document_type

    return "Other"


def load_root(root, source_type):
    if not root.exists():
        return

    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue

        suffix = path.suffix.lower()

        try:
            if suffix == ".pdf":
                pages = extract_pdf(path)
            elif suffix == ".docx":
                pages = extract_docx(path)
            else:
                continue
        except Exception as exc:
            print(f"Warning: could not read {path.name}: {exc}")
            continue

        relative = path.relative_to(root)
        parts = relative.parts

        category = (
            parts[0]
            if len(parts) > 1
            else "Uncategorized"
        )

        all_text = "\n".join(text for _, text in pages)

        DOCS[f"{source_type}:{relative}"] = {
            "name": path.name,
            "relative": str(relative),
            "category": category,
            "year": infer_year(all_text, path.name),
            "document_type": infer_document_type(
                category,
                path.name
            ),
            "source_type": source_type,
            "role": (
                "AUTHORITATIVE / REFERENCE"
                if source_type == "authoritative"
                else "INSTITUTIONAL PRECEDENT / DRAFTING EXAMPLE"
            ),
            "pages": pages,
        }


def load_documents():
    with LOCK:
        DOCS.clear()

        load_root(
            KNOWLEDGE,
            "authoritative"
        )

        load_root(
            INSTITUTIONAL_LIBRARY,
            "institutional_library"
        )


# ------------------------------------------------------------
# Search configuration
# ------------------------------------------------------------
CATEGORY_TERMS = {
    "03_Discipline": [
        "discipline",
        "disciplinary",
        "misconduct",
        "indiscipline",
        "punishment",
        "penalty",
        "suspension",
        "expulsion",
        "appeal",
        "ragging",
        "conduct",
        "offence",
        "offense",
    ],
    "02_Academic Rules": [
        "academic",
        "attendance",
        "registration",
        "semester",
        "course",
        "grade",
        "examination",
        "exam",
        "promotion",
        "degree",
        "undergraduate",
        "ug",
        "b.tech",
        "postgraduate",
        "pg",
        "m.tech",
        "m.sc",
        "ph.d",
        "phd",
        "research scholar",
    ],
    "04_Finance Procurement": [
        "procurement",
        "purchase",
        "tender",
        "gem",
        "gfr",
        "financial",
        "expenditure",
        "bid",
        "quotation",
        "sanction",
        "payment",
        "delegation",
    ],
    "01_Statutes_Gazette": [
        "statute",
        "statutes",
        "gazette",
        "act",
        "board of governors",
        "bog",
        "senate",
        "authority",
        "powers",
        "ordinance",
        "amendment",
        "statutory",
    ],
    "09_Government_Office_Procedure": [
        "csmop",
        "office procedure",
        "file management",
        "noting guidelines",
        "drafting",
    ],
}

LIBRARY_TERMS = {
    "01_Noting": [
        "noting",
        "note",
        "file noting",
        "approval note",
    ],
    "02_Office Orders": [
        "office order",
        "order",
    ],
    "03_Notices": [
        "notice",
    ],
    "04_MoM": [
        "minutes",
        "minute",
        "mom",
        "meeting",
        "proceedings",
    ],
    "05_Minesterial Replies": [
        "ministry",
        "ministerial",
        "lok sabha",
        "rajya sabha",
        "parliamentary",
        "question",
    ],
    "06_Letters": [
        "letter",
        "correspondence",
    ],
    "07_Action Taken Report": [
        "action taken report",
        "atr",
        "action taken",
    ],
    "08_others": [
        "draft",
        "format",
        "precedent",
        "agreement",
        "certificate",
    ],
}

AUTHORITATIVE_CATEGORIES = {
    "01_Statutes_Gazette",
    "02_Academic Rules",
    "03_Discipline",
    "04_Finance Procurement",
    "09_Government_Office_Procedure",
}


def detect_categories(query, mapping):
    query_lower = query.lower()

    scores = {
        category: sum(
            1
            for word in words
            if word in query_lower
        )
        for category, words in mapping.items()
    }

    return sorted(
        [
            category
            for category, score in scores.items()
            if score
        ],
        key=lambda category: (
            -scores[category],
            category,
        ),
    )


def is_rule_question(query):
    query_lower = query.lower()

    markers = [
        "what rule",
        "which rule",
        "applicable rule",
        "regulation",
        "provision",
        "discipline manual",
        "statute",
        "gfr",
        "disciplinary process",
        "competent authority",
        "delegation",
        "procedure",
    ]

    return any(
        marker in query_lower
        for marker in markers
    )


def is_precedent_request(query):
    query_lower = query.lower()

    markers = [
        "previous",
        "past",
        "earlier",
        "historical",
        "precedent",
        "similar",
        "find old",
        "find previous",
        "how has nit sikkim drafted",
    ]

    return any(
        marker in query_lower
        for marker in markers
    )


# ------------------------------------------------------------
# Search
# ------------------------------------------------------------
def make_chunks(text, max_chars=950):
    text = re.sub(r"\s+", " ", text).strip()

    if not text:
        return []

    sentences = re.split(
        r"(?<=[.!?])\s+",
        text
    )

    chunks = []
    current = ""

    for sentence in sentences:
        if (
            len(current)
            + len(sentence)
            + 1
            <= max_chars
        ):
            current = (
                current + " " + sentence
            ).strip()
        else:
            if current:
                chunks.append(current)

            current = sentence

    if current:
        chunks.append(current)

    return chunks


def search_documents(query, limit=10, precedent=False):
    query_lower = query.lower().strip()

    terms = [
        token
        for token in re.findall(
            r"[A-Za-z0-9]+",
            query_lower
        )
        if len(token) >= 3
    ]

    preferred_authority = detect_categories(
        query,
        CATEGORY_TERMS
    )

    preferred_library = detect_categories(
        query,
        LIBRARY_TERMS
    )

    rule_question = is_rule_question(query)

    results = []

    for info in DOCS.values():
        category = info["category"]
        source_type = info["source_type"]
        filename_lower = info["name"].lower()

        for page_number, text in info["pages"]:
            if not text:
                continue

            for chunk in make_chunks(text):
                chunk_lower = chunk.lower()

                score = 0

                # Lexical relevance
                for term in terms:
                    score += (
                        chunk_lower.count(term) * 2
                    )

                    if term in filename_lower:
                        score += 2

                # Category relevance
                if category in preferred_authority:
                    position = preferred_authority.index(
                        category
                    )
                    score += max(
                        35 - position * 8,
                        8
                    )

                if category in preferred_library:
                    position = preferred_library.index(
                        category
                    )
                    score += max(
                        28 - position * 5,
                        5
                    )

                # Exact phrase
                if (
                    query_lower
                    and query_lower in chunk_lower
                ):
                    score += 35

                # Source role
                if rule_question:
                    if source_type == "authoritative":
                        score += 30
                    else:
                        score -= 8

                if precedent:
                    if source_type == "institutional_library":
                        score += 36
                    else:
                        score += 3

                # Strong topical phrases
                for phrase, boost in [
                    ("disciplinary process", 16),
                    ("board of governors", 14),
                    ("office order", 12),
                    ("procurement", 10),
                    ("approval", 5),
                    ("competent authority", 8),
                ]:
                    if (
                        phrase in query_lower
                        and phrase in chunk_lower
                    ):
                        score += boost

                if score > 0:
                    results.append({
                        "score": score,
                        "document": info["name"],
                        "relative": info["relative"],
                        "category": category,
                        "year": info["year"],
                        "document_type": info[
                            "document_type"
                        ],
                        "source_type": source_type,
                        "role": info["role"],
                        "page": page_number,
                        "snippet": chunk,
                    })

    results.sort(
        key=lambda item: (
            -item["score"],
            (
                0
                if (
                    rule_question
                    and item["source_type"]
                    == "authoritative"
                )
                else 1
            ),
            item["document"].lower(),
            item["page"],
        )
    )

    output = []
    seen = set()

    for result in results:
        key = (
            result["document"],
            result["page"],
            result["snippet"][:180],
        )

        if key in seen:
            continue

        seen.add(key)
        output.append(result)

        if len(output) >= limit:
            break

    return output


def build_context(hits):
    parts = []

    for index, hit in enumerate(
        hits,
        start=1
    ):
        parts.append(
            f"[SOURCE {index}]\n"
            f"Document: {hit['document']}\n"
            f"Category: {hit['category']}\n"
            f"Year: {hit['year'] or 'Not identified'}\n"
            f"Type: {hit['document_type']}\n"
            f"Role: {hit['role']}\n"
            f"Page: {hit['page']}\n"
            f"Passage: {hit['snippet']}"
        )

    return "\n\n".join(parts)[
        :MAX_CONTEXT_CHARS
    ]


# ------------------------------------------------------------
# Hermes
# ------------------------------------------------------------
def call_hermes(prompt):
    try:
        result = subprocess.run(
            ["hermes", "-z", prompt],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            cwd=str(BASE),
        )
    except FileNotFoundError:
        return (
            "Hermes was not found in PATH. "
            "Open a new PowerShell window and confirm "
            "`hermes --version` works."
        )
    except subprocess.TimeoutExpired:
        return (
            "Hermes timed out while generating the response."
        )
    except Exception as exc:
        return f"Hermes call failed: {exc}"

    if result.returncode != 0:
        details = (
            result.stderr
            or result.stdout
            or ""
        ).strip()

        return (
            "Hermes returned an error.\n\n"
            + details
        )

    return (
        result.stdout
        or ""
    ).strip()


# ------------------------------------------------------------
# Controlled memory
# ------------------------------------------------------------
def load_memory():
    if not MEMORY_FILE.exists():
        return []

    try:
        data = json.loads(
            MEMORY_FILE.read_text(
                encoding="utf-8"
            )
        )

        return (
            data
            if isinstance(data, list)
            else []
        )
    except Exception:
        return []


def memory_text():
    items = load_memory()

    if not items:
        return "(No approved memory stored.)"

    return "\n".join(
        (
            f"- {item.get('fact', '')} "
            f"| status={item.get('status', '')} "
            f"| confirmed={item.get('confirmed', '')}"
        )
        for item in items
    )


# ------------------------------------------------------------
# Conversational agent
# ------------------------------------------------------------
def chat_with_uri(message, history):
    message = message.strip()

    if not message:
        return (
            "Please tell me what you need."
        ), []

    precedent_request = is_precedent_request(
        message
    )

    hits = search_documents(
        message,
        limit=10,
        precedent=precedent_request,
    )

    recent_history = history[
        -MAX_HISTORY_MESSAGES:
    ]

    history_text = "\n".join(
        (
            f"{item.get('role', 'user').upper()}: "
            f"{item.get('content', '')}"
        )
        for item in recent_history
    )

    prompt = f"""
You are URI, the NIT Sikkim Administrative AI Assistant.

You are conversational-first. The user should be able to describe an
administrative goal naturally. Decide whether to answer, research, draft,
or ask for missing information.

ADMINISTRATIVE ACCURACY RULES:
1. Never invent facts, rules, authorities, dates, penalties, procedures,
   approvals, designations, or recipients.
2. Distinguish CURRENT AUTHORITATIVE / REFERENCE sources from HISTORICAL
   NIT SIKKIM PRECEDENTS / DRAFTING EXAMPLES.
3. Historical examples may show drafting style and past practice, but do not
   automatically establish current authority.
4. Never infer a competent authority from an amount or designation.
5. Distinguish committee recommendation, approval, and final decision.
6. Use stored memory as a convenience, not as a higher authority than
   current authoritative source material.
7. If stored information may be stale, ask the user to confirm it.
8. Ask only essential questions needed to complete the task.
9. If the sources do not establish a point, say so clearly.
10. Do not claim that an action was taken unless the user or an actual tool
    confirms it.

APPROVED MEMORY:
{memory_text()}

RECENT CONVERSATION:
{history_text or "(none)"}

USER MESSAGE:
{message}

LOCAL SOURCES:
{build_context(hits) if hits else "(No relevant local source found.)"}

Respond naturally and concisely.

For a drafting request:
- identify the document/task;
- use the relevant institutional examples;
- ask only essential missing facts before finalizing.

For a rules question:
- answer from authoritative sources where possible;
- identify the source and page;
- do not treat precedent as current rule.

For a historical/precedent request:
- identify the most relevant previous NIT Sikkim examples;
- explain why they are relevant;
- clearly label them as historical precedent.

If a new stable, non-sensitive institutional fact should be remembered, append:
MEMORY SUGGESTION: [fact]

Do NOT suggest memory for:
- passwords or API keys;
- bank/account details;
- allegations;
- disciplinary case details;
- unverified facts;
- temporary facts;
- sensitive personal information.
""".strip()

    return call_hermes(prompt), hits


# ------------------------------------------------------------
# Web UI
# ------------------------------------------------------------
PAGE = r"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>URI - NIT Sikkim Administrative AI Assistant</title>
<style>
body {
  margin: 0;
  background: #f5f7fb;
  color: #1f2937;
  font-family: Segoe UI, Arial, sans-serif;
}
header {
  background: #17324d;
  color: white;
  padding: 18px 24px;
}
header h1 {
  margin: 0;
  font-size: 28px;
}
header p {
  margin: 4px 0 0;
  opacity: .85;
}
.wrap {
  max-width: 1050px;
  margin: 0 auto;
  padding: 18px;
}
.chat {
  height: 70vh;
  overflow-y: auto;
  padding: 6px;
}
.msg {
  display: flex;
  margin: 12px 0;
}
.user {
  justify-content: flex-end;
}
.bubble {
  max-width: 82%;
  padding: 12px 15px;
  border-radius: 14px;
  white-space: pre-wrap;
  line-height: 1.45;
}
.user .bubble {
  background: #17324d;
  color: white;
}
.assistant .bubble {
  background: white;
  border: 1px solid #dbe2ea;
}
.controls {
  display: flex;
  gap: 8px;
  margin-bottom: 8px;
}
.control {
  border: 1px solid #cbd5e1;
  background: white;
  border-radius: 8px;
  padding: 7px 10px;
  cursor: pointer;
}
.composer {
  display: flex;
  gap: 10px;
  background: white;
  border: 1px solid #dbe2ea;
  padding: 10px;
  border-radius: 14px;
}
.composer textarea {
  flex: 1;
  resize: none;
  border: 0;
  outline: 0;
  font: inherit;
  min-height: 54px;
}
.send {
  border: 0;
  background: #17324d;
  color: white;
  padding: 0 20px;
  border-radius: 10px;
  cursor: pointer;
  font-weight: 600;
}
.send:disabled {
  opacity: .6;
}
.status {
  font-size: 12px;
  color: #64748b;
  margin: 8px 4px;
}
.sources {
  margin-top: 10px;
}
.src {
  background: white;
  border: 1px solid #e2e8f0;
  border-radius: 8px;
  padding: 9px;
  margin: 7px 0;
  font-size: 12px;
}
.badge {
  display: inline-block;
  background: #eef2f7;
  border-radius: 999px;
  padding: 3px 8px;
  font-size: 11px;
  margin: 2px 3px 2px 0;
}
</style>
</head>
<body>

<header>
  <h1>URI</h1>
  <p>NIT Sikkim Administrative AI Assistant</p>
</header>

<div class="wrap">

  <div class="controls">
    <button class="control" onclick="startNewChat()">
      New chat
    </button>

    <button class="control" onclick="showApprovedMemory()">
      View approved memory
    </button>
  </div>

  <div id="chat" class="chat"></div>

  <div id="status" class="status">
    Ready.
  </div>

  <div class="composer">
    <textarea
      id="messageInput"
      placeholder="Tell Uri what you need..."
    ></textarea>

    <button
      id="sendButton"
      class="send"
      onclick="sendMessage()"
    >
      Send
    </button>
  </div>

  <div id="sources" class="sources"></div>

</div>

<script>
let chatHistory = [];

const chatBox = document.getElementById("chat");
const messageInput = document.getElementById("messageInput");
const sendButton = document.getElementById("sendButton");
const statusBox = document.getElementById("status");
const sourcesBox = document.getElementById("sources");

function escapeHtml(value) {
  return String(value)
    .replaceAll("&", "&amp;")
    .replaceAll("<", "&lt;")
    .replaceAll(">", "&gt;");
}

function addMessage(role, text) {
  const row = document.createElement("div");
  row.className = "msg " + role;

  const bubble = document.createElement("div");
  bubble.className = "bubble";
  bubble.innerHTML = escapeHtml(text);

  row.appendChild(bubble);
  chatBox.appendChild(row);

  chatBox.scrollTop = chatBox.scrollHeight;
}

function startNewChat() {
  chatHistory = [];
  chatBox.innerHTML = "";
  sourcesBox.innerHTML = "";
  addMessage(
    "assistant",
    "Hello. I am Uri. Tell me what you need."
  );
}

function setBusy(isBusy) {
  sendButton.disabled = isBusy;
  messageInput.disabled = isBusy;

  statusBox.textContent = isBusy
    ? "Uri is working..."
    : "Ready.";
}

function renderSources(sources) {
  if (!sources || !sources.length) {
    sourcesBox.innerHTML = "";
    return;
  }

  const items = sources.map(source => {
    return `
      <div class="src">
        <strong>${escapeHtml(source.document)}</strong>
        - page ${escapeHtml(source.page)}
        <br>
        <span class="badge">
          ${escapeHtml(source.role || "")}
        </span>
        <span class="badge">
          ${escapeHtml(source.document_type || "")}
        </span>
        ${source.year ? `
          <span class="badge">
            ${escapeHtml(source.year)}
          </span>
        ` : ""}
        <br>
        ${escapeHtml(source.snippet)}
      </div>
    `;
  }).join("");

  sourcesBox.innerHTML =
    "<strong>Sources</strong>" + items;
}

async function sendMessage() {
  const text = messageInput.value.trim();

  if (!text) {
    return;
  }

  addMessage("user", text);

  chatHistory.push({
    role: "user",
    content: text
  });

  messageInput.value = "";
  setBusy(true);

  try {
    const response = await fetch(
      "/api/chat",
      {
        method: "POST",
        headers: {
          "Content-Type": "application/json"
        },
        body: JSON.stringify({
          message: text,
          history: chatHistory
        })
      }
    );

    const data = await response.json();

    if (data.error && !data.answer) {
      addMessage(
        "assistant",
        "Error: " + data.error
      );
      return;
    }

    const answer = data.answer || "";

    addMessage(
      "assistant",
      answer
    );

    chatHistory.push({
      role: "assistant",
      content: answer
    });

    renderSources(data.sources || []);

  } catch (error) {
    addMessage(
      "assistant",
      "Connection error: " + error
    );
  } finally {
    setBusy(false);
  }
}

async function showApprovedMemory() {
  try {
    const response = await fetch(
      "/api/memory"
    );

    const data = await response.json();

    if (!data.memory || !data.memory.length) {
      addMessage(
        "assistant",
        "No approved memory is stored yet."
      );
      return;
    }

    const text = data.memory.map(
      item =>
        `${item.fact} [${item.status}; confirmed ${item.confirmed}]`
    ).join("\n");

    addMessage(
      "assistant",
      "Approved memory:\n" + text
    );

  } catch (error) {
    addMessage(
      "assistant",
      "Could not read approved memory."
    );
  }
}

messageInput.addEventListener(
  "keydown",
  function(event) {
    if (
      event.key === "Enter"
      && !event.shiftKey
    ) {
      event.preventDefault();
      sendMessage();
    }
  }
);

startNewChat();
</script>

</body>
</html>
"""


# ------------------------------------------------------------
# HTTP server
# ------------------------------------------------------------
class UriHandler(BaseHTTPRequestHandler):
    def send_response_body(
        self,
        status_code,
        body,
        content_type="text/html; charset=utf-8",
    ):
        encoded = body.encode("utf-8")

        self.send_response(status_code)
        self.send_header(
            "Content-Type",
            content_type,
        )
        self.send_header(
            "Content-Length",
            str(len(encoded)),
        )
        self.end_headers()

        self.wfile.write(encoded)

    def read_json_body(self):
        length = int(
            self.headers.get(
                "Content-Length",
                "0",
            )
        )

        raw = self.rfile.read(length)

        return json.loads(
            raw.decode("utf-8")
        )

    def do_GET(self):
        path = urlparse(
            self.path
        ).path

        if path == "/":
            self.send_response_body(
                200,
                PAGE,
            )
            return

        if path == "/api/health":
            authoritative = sum(
                1
                for info in DOCS.values()
                if info["source_type"]
                == "authoritative"
            )

            institutional = sum(
                1
                for info in DOCS.values()
                if info["source_type"]
                == "institutional_library"
            )

            payload = {
                "ok": True,
                "total_documents": len(DOCS),
                "authoritative_documents": authoritative,
                "institutional_documents": institutional,
                "memory_items": len(
                    load_memory()
                ),
            }

            self.send_response_body(
                200,
                json.dumps(
                    payload
                ),
                "application/json; charset=utf-8",
            )
            return

        if path == "/api/memory":
            self.send_response_body(
                200,
                json.dumps(
                    {
                        "memory": load_memory()
                    },
                    ensure_ascii=False,
                ),
                "application/json; charset=utf-8",
            )
            return

        self.send_response_body(
            404,
            "Not found",
            "text/plain; charset=utf-8",
        )

    def do_POST(self):
        path = urlparse(
            self.path
        ).path

        try:
            data = self.read_json_body()

            if path == "/api/chat":
                message = data.get(
                    "message",
                    "",
                ).strip()

                history = data.get(
                    "history",
                    [],
                )

                answer, sources = chat_with_uri(
                    message,
                    history,
                )

                self.send_response_body(
                    200,
                    json.dumps(
                        {
                            "answer": answer,
                            "sources": sources,
                        },
                        ensure_ascii=False,
                    ),
                    "application/json; charset=utf-8",
                )
                return

            if path == "/api/memory/confirm":
                fact = data.get(
                    "fact",
                    "",
                ).strip()

                if not fact:
                    self.send_response_body(
                        400,
                        json.dumps(
                            {
                                "error":
                                "Memory fact is required."
                            }
                        ),
                        "application/json; charset=utf-8",
                    )
                    return

                memory = load_memory()

                memory.append(
                    {
                        "fact": fact,
                        "status": "CONFIRMED",
                        "confirmed": str(
                            date.today()
                        ),
                    }
                )

                MEMORY_FILE.write_text(
                    json.dumps(
                        memory,
                        ensure_ascii=False,
                        indent=2,
                    ),
                    encoding="utf-8",
                )

                self.send_response_body(
                    200,
                    json.dumps(
                        {"ok": True}
                    ),
                    "application/json; charset=utf-8",
                )
                return

            self.send_response_body(
                404,
                "Not found",
                "text/plain; charset=utf-8",
            )

        except Exception as exc:
            self.send_response_body(
                500,
                json.dumps(
                    {
                        "error": str(exc)
                    },
                    ensure_ascii=False,
                ),
                "application/json; charset=utf-8",
            )


# ------------------------------------------------------------
# Main
# ------------------------------------------------------------
if __name__ == "__main__":
    print(
        "Loading NIT Sikkim knowledge base "
        "and institutional library..."
    )

    load_documents()

    authoritative_count = sum(
        1
        for info in DOCS.values()
        if info["source_type"]
        == "authoritative"
    )

    institutional_count = sum(
        1
        for info in DOCS.values()
        if info["source_type"]
        == "institutional_library"
    )

    print(
        f"Loaded {len(DOCS)} total documents."
    )

    print(
        f"Authoritative/reference documents: "
        f"{authoritative_count}"
    )

    print(
        f"Institutional-library documents: "
        f"{institutional_count}"
    )

    print(
        "URI V5 - Conversational Agent"
    )

    print(
        "Open http://127.0.0.1:8000 "
        "in your browser."
    )

    server = ThreadingHTTPServer(
        ("127.0.0.1", 8000),
        UriHandler,
    )

    server.serve_forever()
