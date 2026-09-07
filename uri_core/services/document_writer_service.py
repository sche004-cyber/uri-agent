import io
import os
try:
    from docx import Document
except ImportError:
    raise ImportError("Please run 'pip install python-docx' in your terminal.")

class DocumentWriterService:
    """
    Takes finalized draft text (Brain-composed - see
    services/document_composer.py) and outputs clean, professionally
    formatted Microsoft Word documents (.docx), parsing markdown tables
    into native Word tables.
    """
    def __init__(self, workspace_path=None):
        if workspace_path is None:
            project_root = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))
            self.workspace_path = os.path.join(project_root, "uri_workspace")
        else:
            self.workspace_path = workspace_path

        os.makedirs(self.workspace_path, exist_ok=True)

    def _build_document(self, draft_text: str) -> Document:
        """The shared rendering core: draft text in, a python-docx
        Document out. Never writes anything itself - both
        save_draft_to_docx (path on disk, legacy callers) and
        render_docx_bytes (in-memory, M19's FileStore-backed path)
        build on this exact same rendering, so there is one document
        renderer, not two."""

        doc = Document()
        lines = draft_text.split("\n")

        table_buffer = []
        in_table = False

        for line in lines:
            stripped = line.strip()

            # Detect markdown table rows
            if stripped.startswith("|") and stripped.endswith("|"):
                # Skip markdown separator rows like |---|---|
                if "---" in stripped:
                    continue

                in_table = True
                cells = [c.strip() for c in stripped.split("|")[1:-1]]
                table_buffer.append(cells)
            else:
                # If we just finished a table block, render it into Word natively
                if in_table:
                    self._add_table_to_doc(doc, table_buffer)
                    table_buffer = []
                    in_table = False

                # Handle normal paragraphs
                if stripped:
                    doc.add_paragraph(stripped)
                else:
                    doc.add_paragraph()

        # Catch any table left at the end of the text
        if in_table and table_buffer:
            self._add_table_to_doc(doc, table_buffer)

        return doc

    def save_draft_to_docx(self, filename: str, draft_text: str) -> str:
        """Parses draft text and builds native Word tables where markdown tables appear."""
        file_path = os.path.join(self.workspace_path, filename)

        try:
            doc = self._build_document(draft_text)
            doc.save(file_path)
            return file_path
        except Exception as e:
            return f"[ERROR] Could not save document: {str(e)}"

    def render_docx_bytes(self, draft_text: str) -> bytes:
        """M19: the same rendering as save_draft_to_docx, but returned
        as in-memory bytes rather than written to a fixed path - so the
        caller (generate_document.py) can hand them to FileStore
        (containment-checked, size/type-validated, session-scoped and
        downloadable via the existing GET /files/{id}/content) instead
        of writing outside that boundary. Raises on failure rather than
        returning an "[ERROR] ..." string, so a caller cannot mistake a
        failure message for real document bytes."""

        doc = self._build_document(draft_text)
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def _add_table_to_doc(self, doc, table_data):
        """Helper method to construct native Word tables with clean borders."""
        if not table_data:
            return
            
        num_rows = len(table_data)
        num_cols = max(len(row) for row in table_data)
        
        table = doc.add_table(rows=num_rows, cols=num_cols)
        table.style = 'Table Grid'
        
        for r_idx, row_values in enumerate(table_data):
            for c_idx, val in enumerate(row_values):
                if c_idx < num_cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = val
        
        # Add spacing after table
        doc.add_paragraph()
