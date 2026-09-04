from pathlib import Path
from openpyxl import load_workbook


class StudentRecordService:

    HEADER_ALIASES = {
        "roll_number": [
            "roll no",
            "roll no.",
            "enrollment no",
            "enrolment no",
            "enrollement no",
            "enrolment no allotted",
            "enrollment no.",
        ],
        "name": [
            "name",
            "candidate name",
            "name of student",
            "name of the student",
            "name of applicant",
            "name of the applicant",
        ],
    }

    def __init__(
        self,
        records_path="uri_workspace/linked_documents/Insurance_extracted"
    ):
        self.records_path = Path(records_path)

    def search_student(self, roll_number):
        query = str(roll_number).strip().upper()
        matches = []

        if not self.records_path.exists():
            return {
                "success": False,
                "error": f"Records path not found: {self.records_path}",
                "matches": []
            }

        files = [
            file for file in self.records_path.glob("*.xlsx")
            if not file.name.startswith("~$")
        ]

        for file_path in files:
            try:
                workbook = load_workbook(
                    file_path,
                    read_only=True,
                    data_only=True
                )

                for worksheet in workbook.worksheets:
                    match = self._search_worksheet(
                        worksheet,
                        query,
                        file_path.name
                    )

                    if match:
                        matches.extend(match)

                workbook.close()

            except Exception as error:
                matches.append({
                    "source_file": file_path.name,
                    "error": str(error)
                })

        return {
            "success": True,
            "query": query,
            "match_count": len(matches),
            "matches": matches
        }

    def _search_worksheet(
        self,
        worksheet,
        query,
        source_file
    ):
        rows = list(
            worksheet.iter_rows(
                min_row=1,
                max_row=10,
                values_only=True
            )
        )

        header_row_number = None
        headers = None
        roll_column_index = None

        for row_number, row in enumerate(rows, start=1):

            normalized = [
                self._normalize_header(value)
                for value in row
            ]

            for index, value in enumerate(normalized):
                if value in self.HEADER_ALIASES["roll_number"]:
                    header_row_number = row_number
                    headers = list(row)
                    roll_column_index = index
                    break

            if roll_column_index is not None:
                break

        if roll_column_index is None:
            return []

        name_column_index = self._find_name_column(headers)

        matches = []

        for row in worksheet.iter_rows(
            min_row=header_row_number + 1,
            values_only=True
        ):

            if roll_column_index >= len(row):
                continue

            roll_value = row[roll_column_index]

            if roll_value is None:
                continue

            if str(roll_value).strip().upper() != query:
                continue

            record = {
                "roll_number": str(roll_value).strip(),
                "source_file": source_file,
                "sheet": worksheet.title,
                "header_row": header_row_number,
                "record": {}
            }

            if (
                name_column_index is not None
                and name_column_index < len(row)
            ):
                record["student_name"] = (
                    str(row[name_column_index]).strip()
                    if row[name_column_index] is not None
                    else None
                )

            for index, header in enumerate(headers):

                if header is None:
                    continue

                value = (
                    row[index]
                    if index < len(row)
                    else None
                )

                record["record"][str(header).strip()] = value

            matches.append(record)

        return matches

    def _find_name_column(self, headers):

        for index, header in enumerate(headers):

            normalized = self._normalize_header(header)

            if normalized in self.HEADER_ALIASES["name"]:
                return index

        return None

    @staticmethod
    def _normalize_header(value):

        if value is None:
            return ""

        return (
            str(value)
            .strip()
            .lower()
            .replace("_", " ")
            .replace("  ", " ")
        )
