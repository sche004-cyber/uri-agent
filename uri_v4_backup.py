import json
import re
import subprocess
import threading
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse

# ============================================================
# URI V3 — NIT Sikkim Administrative AI Assistant
# Administrative Accuracy Mode
#
# Local knowledge retrieval + Hermes AI
# ============================================================

BASE = Path(__file__).resolve().parent
KNOWLEDGE = BASE / "knowledge"
DOCS = {}
LOCK = threading.Lock()

MAX_CONTEXT_CHARS = 12000


# ------------------------------------------------------------
# Document extraction
# ------------------------------------------------------------
def extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []

    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append((i, text))

    return pages


def extract_docx(path):
    from docx import Document

    doc = Document(str(path))
    lines = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            values = [value for value in values if value]
            if values:
                lines.append(" | ".join(values))

    return [(1, "\n".join(lines))]


def load_docs():
    with LOCK:
        DOCS.clear()

        for path in sorted(KNOWLEDGE.rglob("*")):
            if not path.is_file():
                continue

            suffix = path.suffix.lower()

            try:
                if suffix == ".pdf":
                    pages = extract_pdf(path)
                elif suffix == ".docx":
                    pages = extract_docx(path)
                else:
                    continue
            except Exception as exc:
                print(f"Warning: could not read {path.name}: {exc}")
                continue

            relative = path.relative_to(KNOWLEDGE)
            category = relative.parts[0] if len(relative.parts) > 1 else "Uncategorized"

            DOCS[str(relative)] = {
                "name": path.name,
                "category": category,
                "pages": pages,
            }


# ------------------------------------------------------------
# Chunking
# ------------------------------------------------------------
def sentence_chunks(text, max_chars=900):
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []

    sentences = re.split(r"(?<=[.!?])\s+", text)
    output = []
    current = ""

    for sentence in sentences:
        if len(current) + len(sentence) + 1 <= max_chars:
            current = (current + " " + sentence).strip()
        else:
            if current:
                output.append(current)
            current = sentence

    if current:
        output.append(current)

    return output


# ------------------------------------------------------------
# Category detection and source hierarchy
# ------------------------------------------------------------
CATEGORY_TERMS = {
    "03_Discipline": [
        "discipline", "disciplinary", "misconduct", "indiscipline",
        "punishment", "penalty", "hostel", "suspension", "expulsion",
        "appeal", "ragging", "conduct", "offence", "offense",
        "disciplinary action", "student misconduct",
    ],
    "02_Academic Rules": [
        "academic", "attendance", "registration", "semester", "course",
        "grade", "examination", "exam", "promotion", "degree",
        "undergraduate", "ug", "b.tech", "postgraduate", "pg",
        "m.tech", "m.sc", "ph.d", "phd", "research scholar",
    ],
    "04_Finance Procurement": [
        "procurement", "purchase", "tender", "gem", "gfr",
        "financial", "expenditure", "bid", "quotation", "sanction",
        "purchase committee", "payment", "delegation of financial powers",
    ],
    "01_Statutes_Gazette": [
        "statute", "statutes", "gazette", "act", "board of governors",
        "bog", "senate", "authority", "powers", "ordinance",
        "amendment", "statutory",
    ],
    "05_Office Order": [
        "office order", "order format", "office order format",
    ],
    "06_Noting": [
        "noting", "note", "file noting", "approval note", "proposal",
        "put up", "submitted for approval",
    ],
    "07_Minute": [
        "minutes", "minute", "mom", "meeting", "proceedings",
        "deliberation", "agenda",
    ],
    "08_Institutional Information": [
        "annual report", "institute", "department", "faculty",
        "staff", "campus", "institution", "vision", "mission",
    ],
}

# Stronger document-role precedence:
# authoritative rules should outrank institutional examples for rule questions.
AUTHORITATIVE_CATEGORIES = {
    "01_Statutes_Gazette",
    "02_Academic Rules",
    "03_Discipline",
    "04_Finance Procurement",
}

EXAMPLE_CATEGORIES = {
    "05_Office Order",
    "06_Noting",
    "07_Minute",
}

LOWER_PRIORITY_CATEGORIES = {
    "08_Institutional Information",
}


def detect_categories(query):
    q = query.lower()
    scores = {}

    for category, words in CATEGORY_TERMS.items():
        score = sum(1 for word in words if word in q)
        if score:
            scores[category] = score

    return sorted(scores, key=lambda key: (-scores[key], key))


def search_docs(query, limit=8):
    qlower = query.lower().strip()

    terms = [
        token.lower()
        for token in re.findall(r"[A-Za-z0-9]+", query)
        if len(token) >= 3
    ]

    preferred = detect_categories(query)
    is_rule_question = any(
        marker in qlower
        for marker in [
            "what rule", "which rule", "applicable rule", "rules",
            "regulation", "provision", "discipline manual",
            "statute", "gfr", "disciplinary process", "punishment",
            "authority", "competent authority",
        ]
    )

    scored = []

    for doc_key, info in DOCS.items():
        name = info["name"]
        category = info["category"]

        if category in preferred:
            category_rank = preferred.index(category)
        else:
            category_rank = 99

        for page_no, text in info["pages"]:
            if not text:
                continue

            for chunk in sentence_chunks(text):
                cl = chunk.lower()
                score = 0

                # Query term relevance.
                for term in terms:
                    count = cl.count(term)
                    score += count * 2
                    if term in name.lower():
                        score += 1

                # Strong category preference.
                if category in preferred:
                    score += max(30 - category_rank * 8, 8)

                # For rule questions, authoritative rule categories get
                # a strong boost; examples and annual-report material are
                # intentionally down-ranked.
                if is_rule_question:
                    if category in AUTHORITATIVE_CATEGORIES:
                        score += 20
                    elif category in EXAMPLE_CATEGORIES:
                        score -= 8
                    elif category in LOWER_PRIORITY_CATEGORIES:
                        score -= 15

                # Exact phrase and section relevance.
                if qlower and qlower in cl:
                    score += 30

                if "disciplinary process" in qlower and "disciplinary process" in cl:
                    score += 25

                if "appeal" in qlower and "appeal" in cl:
                    score += 8

                if "procurement" in qlower and "procurement" in cl:
                    score += 8

                if score > 0:
                    scored.append(
                        {
                            "score": score,
                            "document": name,
                            "category": category,
                            "page": page_no,
                            "snippet": chunk,
                        }
                    )

    scored.sort(
        key=lambda item: (
            -item["score"],
            0 if item["category"] in AUTHORITATIVE_CATEGORIES else 1,
            item["document"].lower(),
            item["page"],
        )
    )

    output = []
    seen = set()

    for hit in scored:
        key = (
            hit["document"],
            hit["page"],
            hit["snippet"][:180],
        )

        if key in seen:
            continue

        seen.add(key)
        output.append(hit)

        if len(output) >= limit:
            break

    return output


# ------------------------------------------------------------
# Hermes integration
# ------------------------------------------------------------
def call_hermes(prompt):
    try:
        result = subprocess.run(
            ["hermes", "-z", prompt],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            cwd=str(BASE),
        )
    except FileNotFoundError:
        return (
            "Hermes could not be found in PATH. Open a new PowerShell "
            "window and confirm that `hermes --version` works."
        )
    except subprocess.TimeoutExpired:
        return "Hermes timed out while generating the response."
    except Exception as exc:
        return f"Hermes call failed: {exc}"

    if result.returncode != 0:
        details = (result.stderr or result.stdout or "").strip()
        return f"Hermes returned an error.\n\n{details}"

    return (result.stdout or "").strip()


def build_context(hits):
    parts = []

    for index, hit in enumerate(hits, start=1):
        parts.append(
            f"[SOURCE {index}]\n"
            f"Document: {hit['document']}\n"
            f"Category: {hit['category']}\n"
            f"Page: {hit['page']}\n"
            f"Passage: {hit['snippet']}"
        )

    context = "\n\n".join(parts)
    return context[:MAX_CONTEXT_CHARS]


def answer_from_sources(query, hits):
    if not hits:
        return (
            "I could not locate a relevant passage in the local "
            "NIT Sikkim knowledge base.\n\n"
            "VERIFICATION / APPROVAL POINTS:\n"
            "- Review the complete knowledge base or provide the relevant document."
        )

    context = build_context(hits)

    prompt = f"""
You are URI, the NIT Sikkim Administrative AI Assistant.

ADMINISTRATIVE ACCURACY MODE

Answer the user's question using the supplied source passages.

NON-NEGOTIABLE RULES:
1. Never invent a rule, section, clause, authority, power, penalty,
   date, fact, procedure, or institutional practice.
2. Do not add facts merely because they sound plausible.
3. If a point is not established by the supplied sources, say:
   "The supplied sources do not establish this point."
4. Distinguish authoritative rules/manuals/statutes from Office Orders,
   Notings, Minutes, Annual Reports, and other examples.
5. A prior institutional document is not automatically an authority for
   what is legally or administratively permissible.
6. Distinguish allegations, committee observations/findings,
   recommendations, approvals, and final decisions.
7. Do not infer a competent authority from an amount, designation,
   or common practice unless the supplied sources establish it.
8. For disciplinary matters, distinguish suggested measures from mandatory
   provisions and distinguish recommendations from final punishment.
9. Do not use a source merely because a keyword matches. Use it only if
   it actually supports the answer.
10. Be concise and formal.

OUTPUT:
Answer:
[direct answer]

Source(s):
[only the sources actually relied upon, with document and page]

Verification / Approval Points:
[only matters that genuinely require human verification]
If none, state "None identified from the supplied sources."

USER QUESTION:
{query}

SUPPLIED SOURCES:
{context}
""".strip()

    return call_hermes(prompt)


# ------------------------------------------------------------
# Drafting helpers
# ------------------------------------------------------------
def draft_noting(data):
    subject = data.get("subject", "").strip() or "[Subject]"
    background = (
        data.get("background", "").strip()
        or "[Background / reference to be inserted]"
    )
    requirement = (
        data.get("requirement", "").strip()
        or "[Requirement / proposal details]"
    )
    cost = data.get("cost", "").strip() or "[Amount, if applicable]"
    authority = (
        data.get("authority", "").strip()
        or "[Competent authority — TO BE VERIFIED]"
    )
    extra = data.get("extra", "").strip()

    query = " ".join(
        [subject, background, requirement, "noting", "proposal", "approval"]
    )

    hits = search_docs(query, 8)
    context = build_context(hits)

    prompt = f"""
You are URI, preparing a concise official administrative noting for
National Institute of Technology Sikkim.

ADMINISTRATIVE ACCURACY MODE

Prepare an office-ready noting using ONLY the facts supplied below,
together with institutional procedures actually supported by the supplied
source passages.

ABSOLUTE RULES:
1. Never invent factual details.
2. Never add justification that the officer did not supply unless a source
   explicitly supports it as a fact.
3. Never assume the competent authority from the amount or designation.
4. Never create a budget head, sanction, finance concurrence, procurement
   route, approval power, or delegation that is not established.
5. Do not use an Annual Report as the primary legal/procedural authority
   when a dedicated rule/manual source is available.
6. Previous Office Orders, Notings and MoMs may be used to understand
   drafting style and institutional precedent, but not as automatic proof
   of current authority.
7. Where information is missing, write "To be verified" or "To be confirmed"
   rather than guessing.
8. Keep the noting concise and formal.

STRUCTURE:
Subject
1. Background / Reference
2. Requirement / Facts
3. Applicable Rule / Procedure
4. Financial Implication
5. Proposal
6. Submitted for approval/orders

Then:
VERIFICATION / APPROVAL POINTS

Only include a rule when the supplied source actually supports it.

OFFICER-SUPPLIED INFORMATION:
Subject: {subject}
Background / Reference: {background}
Requirement / Facts: {requirement}
Estimated Cost: {cost}
Approval Level Provided by Officer: {authority}
Additional Information: {extra}

SOURCE PASSAGES:
{context}
""".strip()

    return call_hermes(prompt), hits


def draft_office_order(instruction):
    instruction = instruction.strip() or "[Decision / direction to be issued]"

    hits = search_docs(instruction + " office order approval", 8)
    context = build_context(hits)

    prompt = f"""
You are URI, preparing a draft Office Order for NIT Sikkim.

ADMINISTRATIVE ACCURACY MODE

Draft the order using the supplied instruction and source passages.

RULES:
1. Do not invent facts, dates, reference numbers, names, designations,
   approvals, authorities or financial powers.
2. Use placeholders for missing reference/date/details.
3. A committee recommendation must not be presented as a final decision
   unless the sources establish approval/finality.
4. Previous Office Orders provide formatting/style examples, not automatic
   authority for a current decision.
5. Keep the wording formal, concise and institutional.

Use this general structure where applicable:
OFFICE ORDER
[operative decision/direction]
[approval statement only when supported]
[signing designation]
Copy of Information to:

After the draft, include:
VERIFICATION / APPROVAL POINTS

INSTRUCTION:
{instruction}

SOURCE PASSAGES:
{context}
""".strip()

    return call_hermes(prompt), hits


# ------------------------------------------------------------
# Web interface
# ------------------------------------------------------------
PAGE = r"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>URI — NIT Sikkim Administrative AI Assistant</title>
<style>
body{
    font-family:Segoe UI,Arial,sans-serif;
    background:#f5f7fb;margin:0;color:#1f2937
}
header{
    background:#17324d;color:#fff;padding:22px 28px
}
header h1{margin:0;font-size:30px}
header p{margin:5px 0 0;opacity:.85}
.wrap{max-width:1150px;margin:24px auto;padding:0 18px}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:18px}
.card{
    background:#fff;border:1px solid #dbe2ea;border-radius:12px;
    padding:18px;box-shadow:0 2px 10px rgba(0,0,0,.04)
}
h2{margin-top:0}
.tabs button{
    border:0;background:#e9eef5;padding:10px 14px;border-radius:8px;
    margin:0 6px 6px 0;cursor:pointer
}
.tabs button.active{background:#17324d;color:#fff}
textarea,input{
    width:100%;box-sizing:border-box;padding:10px;
    border:1px solid #cbd5e1;border-radius:8px;margin:6px 0 12px;
    font:inherit
}
button.primary{
    background:#17324d;color:#fff;border:0;padding:11px 16px;
    border-radius:8px;cursor:pointer
}
button.primary:disabled{opacity:.6}
pre{
    white-space:pre-wrap;background:#f8fafc;border:1px solid #e2e8f0;
    padding:14px;border-radius:8px;max-height:650px;overflow:auto
}
.src{
    margin:8px 0;padding:10px;background:#f8fafc;
    border-left:4px solid #17324d
}
.muted{color:#64748b;font-size:13px}
.status{font-size:13px;margin-top:8px}
</style>
</head>
<body>
<header>
<h1>URI</h1>
<p>NIT Sikkim Administrative AI Assistant — V3 Administrative Accuracy Mode</p>
</header>

<div class="wrap">
<div class="card">
<div class="tabs">
<button class="active" onclick="show('noting',this)">Prepare Noting</button>
<button onclick="show('rules',this)">Ask Rules</button>
<button onclick="show('office',this)">Draft Office Order</button>
</div>

<div id="noting">
<h2>Prepare Noting</h2>
<input id="subject" placeholder="Subject">
<textarea id="background" rows="4" placeholder="Background / reference"></textarea>
<textarea id="requirement" rows="4" placeholder="Requirement / proposal details"></textarea>
<input id="cost" placeholder="Estimated cost, e.g. ₹8,00,000">
<input id="authority" placeholder="Approval level (optional)">
<textarea id="extra" rows="3" placeholder="Additional points (optional)"></textarea>
<button id="notingBtn" class="primary" onclick="draftNoting()">Generate Noting</button>
</div>

<div id="rules" style="display:none">
<h2>Ask Rules</h2>
<textarea id="query" rows="5" placeholder="Example: What does the Discipline Manual say about the disciplinary process?"></textarea>
<button id="rulesBtn" class="primary" onclick="askRules()">Ask Uri</button>
</div>

<div id="office" style="display:none">
<h2>Draft Office Order</h2>
<textarea id="oo" rows="8" placeholder="Describe the decision/direction to be issued."></textarea>
<button id="officeBtn" class="primary" onclick="draftOO()">Generate Office Order</button>
</div>
</div>

<div class="grid" style="margin-top:18px">
<div class="card">
<h2>Output</h2>
<pre id="output">Uri is ready.</pre>
<div id="status" class="status muted"></div>
</div>
<div class="card">
<h2>Sources</h2>
<div id="sources" class="muted">Sources used for rule-aware tasks will appear here.</div>
</div>
</div>
</div>

<script>
function show(id,button){
    for(const name of ['noting','rules','office']){
        document.getElementById(name).style.display=(name===id?'block':'none');
    }
    document.querySelectorAll('.tabs button').forEach(x=>x.classList.remove('active'));
    button.classList.add('active');
}

function setBusy(id,busy){
    const btn=document.getElementById(id);
    if(!btn.dataset.normal) btn.dataset.normal=btn.textContent;
    btn.disabled=busy;
    btn.textContent=busy?'Working...':btn.dataset.normal;
}

function escapeHtml(value){
    return String(value)
      .replaceAll('&','&amp;')
      .replaceAll('<','&lt;')
      .replaceAll('>','&gt;')
      .replaceAll('"','&quot;')
      .replaceAll("'","&#039;");
}

function renderSources(sources){
    const el=document.getElementById('sources');

    if(!sources || !sources.length){
        el.innerHTML='<span class="muted">No source passages located.</span>';
        return;
    }

    el.innerHTML=sources.map(s=>`
        <div class="src">
            <strong>${escapeHtml(s.document)}</strong><br>
            ${escapeHtml('Page ' + s.page)}<br>
            <span class="muted">${escapeHtml(s.snippet)}</span>
        </div>`).join('');
}

async function askRules(){
    const query=document.getElementById('query').value.trim();
    if(!query) return;

    setBusy('rulesBtn',true);
    document.getElementById('status').textContent=
        'Searching local sources and asking Hermes...';

    try{
        const r=await fetch('/api/rules',{
            method:'POST',
            headers:{'Content-Type':'application/json'},
            body:JSON.stringify({query})
        });

        const data=await r.json();
        document.getElementById('output').textContent=data.answer || data.error || '';
        renderSources(data.sources || []);
        document.getElementById('status').textContent='Completed.';
    }catch(e){
        document.getElementById('output').textContent='Error: '+e;
        document.getElementById('status').textContent='';
    }finally{
        setBusy('rulesBtn',false);
    }
}

async function draftNoting(){
    const payload={
        subject:document.getElementById('subject').value,
        background:document.getElementById('background').value,
        requirement:document.getElementById('requirement').value,
        cost:document.getElementById('cost').value,
        authority:document.getElementById('authority').value,
        extra:document.getElementById('extra').value
    };

    setBusy('notingBtn',true);
    document.getElementById('status').textContent=
        'Preparing an accuracy-controlled noting with Hermes...';

    try{
        const r=await fetch('/api/noting',{
            method:'POST',
            headers:{'Content-Type':'application/json'},
            body:JSON.stringify(payload)
        });

        const data=await r.json();
        document.getElementById('output').textContent=data.answer || data.error || '';
        renderSources(data.sources || []);
        document.getElementById('status').textContent='Completed.';
    }catch(e){
        document.getElementById('output').textContent='Error: '+e;
        document.getElementById('status').textContent='';
    }finally{
        setBusy('notingBtn',false);
    }
}

async function draftOO(){
    const instruction=document.getElementById('oo').value.trim();
    if(!instruction) return;

    setBusy('officeBtn',true);
    document.getElementById('status').textContent=
        'Preparing accuracy-controlled Office Order with Hermes...';

    try{
        const r=await fetch('/api/office-order',{
            method:'POST',
            headers:{'Content-Type':'application/json'},
            body:JSON.stringify({instruction})
        });

        const data=await r.json();
        document.getElementById('output').textContent=data.answer || data.error || '';
        renderSources(data.sources || []);
        document.getElementById('status').textContent='Completed.';
    }catch(e){
        document.getElementById('output').textContent='Error: '+e;
        document.getElementById('status').textContent='';
    }finally{
        setBusy('officeBtn',false);
    }
}
</script>
</body>
</html>
"""


# ------------------------------------------------------------
# HTTP server
# ------------------------------------------------------------
class Handler(BaseHTTPRequestHandler):
    def _send(self, code, body, content_type="text/html; charset=utf-8"):
        encoded = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)

    def _json_body(self):
        length = int(self.headers.get("Content-Length", "0"))
        raw = self.rfile.read(length)
        return json.loads(raw.decode("utf-8"))

    def do_GET(self):
        path = urlparse(self.path).path

        if path == "/":
            self._send(200, PAGE)
            return

        if path == "/api/health":
            self._send(
                200,
                json.dumps(
                    {
                        "ok": True,
                        "documents": len(DOCS),
                        "mode": "Administrative Accuracy Mode",
                    }
                ),
                "application/json; charset=utf-8",
            )
            return

        self._send(404, "Not found", "text/plain; charset=utf-8")

    def do_POST(self):
        path = urlparse(self.path).path

        try:
            data = self._json_body()

            if path == "/api/rules":
                query = data.get("query", "").strip()
                hits = search_docs(query, 8)
                answer = answer_from_sources(query, hits)

                self._send(
                    200,
                    json.dumps(
                        {
                            "answer": answer,
                            "sources": hits,
                        },
                        ensure_ascii=False,
                    ),
                    "application/json; charset=utf-8",
                )
                return

            if path == "/api/noting":
                answer, hits = draft_noting(data)

                self._send(
                    200,
                    json.dumps(
                        {
                            "answer": answer,
                            "sources": hits,
                        },
                        ensure_ascii=False,
                    ),
                    "application/json; charset=utf-8",
                )
                return

            if path == "/api/office-order":
                instruction = data.get("instruction", "").strip()
                answer, hits = draft_office_order(instruction)

                self._send(
                    200,
                    json.dumps(
                        {
                            "answer": answer,
                            "sources": hits,
                        },
                        ensure_ascii=False,
                    ),
                    "application/json; charset=utf-8",
                )
                return

            self._send(404, "Not found", "text/plain; charset=utf-8")

        except Exception as exc:
            self._send(
                500,
                json.dumps({"error": str(exc)}, ensure_ascii=False),
                "application/json; charset=utf-8",
            )


if __name__ == "__main__":
    print("Loading NIT Sikkim knowledge base...")
    load_docs()
    print(f"Loaded {len(DOCS)} documents.")
    print("URI V3 — Administrative Accuracy Mode")
    print("Open http://127.0.0.1:8000 in your browser.")

    server = ThreadingHTTPServer(("127.0.0.1", 8000), Handler)
    server.serve_forever()
